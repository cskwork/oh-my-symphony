#!/usr/bin/env python3
"""Installed-wheel PM Workspace lifecycle proof on a disposable Git project."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

from symphony._shell import process_group_exists, process_identity
from symphony.orchestrator.run_registry import RunRegistry
from symphony.pm.authority import PmAuthority
from symphony.pm.models import PmError
from symphony.trackers.file import FileBoardTracker
from symphony.workflow import WorkflowState
import symphony.pm


def run(*args: str, cwd: Path) -> str:
    return subprocess.check_output(args, cwd=cwd, text=True).strip()


def write(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8")


def stop_process(process: subprocess.Popen[str]) -> None:
    process.terminate()
    try:
        process.wait(timeout=5)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


def wait_for_api(api: "Api", process: subprocess.Popen[str], log_path: Path) -> None:
    for _ in range(100):
        try:
            api.call("GET", "/capabilities")
            return
        except Exception:
            if process.poll() is not None:
                raise RuntimeError(log_path.read_text(encoding="utf-8"))
            time.sleep(0.1)
    raise RuntimeError("server did not become ready")


def remove_planning_marker(path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines(keepends=True)
    output: list[str] = []
    skipping = False
    removed = False
    for line in lines:
        if not skipping and line.rstrip("\r\n") == "symphony_pm:":
            skipping = True
            removed = True
            continue
        if skipping and (line.startswith(" ") or line.strip() == ""):
            continue
        skipping = False
        output.append(line)
    if not removed:
        raise AssertionError(f"PM marker was not present in {path}")
    path.write_text("".join(output), encoding="utf-8")


def set_fixture_paused(
    project_root: Path,
    ticket_identifiers: tuple[str, ...],
    *,
    paused: bool,
) -> None:
    registry = RunRegistry(project_root / ".symphony" / "state.db")
    try:
        for identifier in ticket_identifiers:
            registry.set_issue_flags(
                identifier,
                paused=paused,
                pause_reason=(
                    "installed journey evidence hold" if paused else None
                ),
            )
            flags = registry.get_issue_flags(identifier)
            if paused:
                assert flags is not None and flags.paused
                assert flags.pause_reason == "installed journey evidence hold"
            else:
                assert flags is None or (
                    not flags.paused and flags.pause_reason is None
                )
    finally:
        registry.close()


def read_ticket_states(
    project_root: Path, ticket_identifiers: tuple[str, ...]
) -> dict[str, str]:
    state = WorkflowState(project_root / "WORKFLOW.md")
    config, error = state.reload()
    assert config is not None and error is None, error
    tracker = FileBoardTracker(config.tracker)
    result: dict[str, str] = {}
    for identifier in ticket_identifiers:
        issue = tracker.fetch_issue_full_by_id(identifier)
        assert issue is not None
        result[identifier] = issue.state
    return result


def exercise_fallback_authority(
    project_root: Path,
    project_id: str,
    ticket_identifier: str,
) -> dict[str, object]:
    database = project_root / ".symphony" / "state.db"
    fallback = database.with_name(database.name + ".pm-managed.json")
    assert database.is_file() and fallback.is_file()
    originals = tuple(
        path for path in (
            database,
            database.with_name(database.name + "-wal"),
            database.with_name(database.name + "-shm"),
        ) if path.exists()
    )
    original_hashes = {path.name: sha256(path) for path in originals}
    backups = {
        path: path.with_name(path.name + ".combined-backup")
        for path in originals
    }
    assert not any(path.exists() for path in backups.values())
    for path, backup in backups.items():
        os.replace(path, backup)
    temporary: RunRegistry | None = None
    try:
        temporary = RunRegistry(database)
        state = WorkflowState(project_root / "WORKFLOW.md")
        config, error = state.reload()
        assert config is not None and error is None, error
        issue = FileBoardTracker(config.tracker).fetch_issue_full_by_id(
            ticket_identifier
        )
        assert issue is not None and issue.planning_ref is None
        membership = PmAuthority(
            temporary, project_root / ".symphony" / "pm-locks"
        ).recognize(issue, project_id)
        assert membership.managed and not membership.runnable
        assert "fallback" in (membership.reason or "")
        return {
            "recognized": True,
            "runnable": False,
            "reason": membership.reason,
        }
    finally:
        if temporary is not None:
            temporary.close()
        database.unlink(missing_ok=True)
        for suffix in ("-wal", "-shm"):
            database.with_name(database.name + suffix).unlink(missing_ok=True)
        for path, backup in backups.items():
            os.replace(backup, path)
        assert {
            path.name: sha256(path) for path in originals
        } == original_hashes


class Api:
    def __init__(self, port: int) -> None:
        self.base = f"http://127.0.0.1:{port}/api/v1/pm"

    def call(self, method: str, path: str, body: dict | None = None) -> dict:
        data = json.dumps(body).encode() if body is not None else None
        request = urllib.request.Request(
            self.base + path, data=data, method=method,
            headers={
                "Content-Type": "application/json",
                "X-Symphony-PM-Owner": "local-operator",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=240) as response:
                return json.loads(response.read() or b"{}")
        except urllib.error.HTTPError as exc:
            payload = exc.read().decode(errors="replace")
            raise RuntimeError(f"{method} {path}: HTTP {exc.code}: {payload}") from exc

    def read(self, path: str) -> bytes:
        request = urllib.request.Request(
            self.base + path,
            headers={"X-Symphony-PM-Owner": "local-operator"},
        )
        with urllib.request.urlopen(request, timeout=60) as response:
            return response.read()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def distribution_inventory() -> dict[str, str]:
    names = ("oh-my-symphony", "oh-my-symphony-pm", "openpyxl", "defusedxml", "et-xmlfile")
    return {
        name: importlib.metadata.version(name)
        for name in names
        if _distribution_exists(name)
    }


def _distribution_exists(name: str) -> bool:
    try:
        importlib.metadata.distribution(name)
    except importlib.metadata.PackageNotFoundError:
        return False
    return True


def envelope(project: str, change: str, revision: str | None = None, payload: dict | None = None, operation: str | None = None) -> dict:
    return {"project_id": project, "change_id": change, "expected_revision": revision, "client_operation_id": operation, "payload": payload or {}}


def setup(
    root: Path, openspec: Path, port: int, *, phase_reach: bool = False,
) -> tuple[Path, Path, Path]:
    project = root / "project"
    planning = root / "planning"
    project.mkdir(parents=True)
    run("git", "init", "-q", cwd=project)
    run("git", "config", "user.name", "PM E2E", cwd=project)
    run("git", "config", "user.email", "pm-e2e@example.invalid", cwd=project)
    (project / "kanban").mkdir()
    schema_resource = Path(symphony.pm.__file__).parent / "resources" / "openspec" / "symphony-pm"
    schema_root = project / "openspec" / "schemas" / "symphony-pm"
    shutil.copytree(schema_resource, schema_root)
    write(project / "openspec" / "config.yaml", "schema: spec-driven\n")
    write(project / "openspec" / "specs" / ".gitkeep", "")
    write(
        project / "openspec" / "specs" / "requirements" / "spec.md",
        """# Existing accepted baseline specification

## Purpose
Preserve accepted behavior while an explicit planning change is reviewed.

## Requirements

### Requirement: Existing accepted baseline
<!-- symphony:req=REQ-A1B2C3D4 -->

The existing system SHALL retain accepted behavior while a change is reviewed.

#### Scenario: Baseline remains available
<!-- symphony:scenario=SC-A1B2C3D5 -->
- **WHEN** a proposed change is opened
- **THEN** the accepted baseline remains visible for comparison
""",
    )
    write(project / "openspec" / "changes" / "archive" / ".gitkeep", "")
    workflow = project / "WORKFLOW.md"
    write(workflow, f"""---
tracker:
  kind: file
  board_root: ./kanban
  project_slug: pm-installed-e2e
  active_states: [Todo]
  terminal_states: [Done, Cancelled]
polling:
  interval_ms: 500
workspace:
  root: {root / "worker-workspaces"}
  reuse_policy: preserve
hooks:
  after_create: |
    git -C "$SYMPHONY_WORKFLOW_DIR" worktree add --force \
      -b "symphony/$(basename "$PWD")" "$PWD" HEAD
  after_done: exit 17
agent:
  kind: codex
  max_concurrent_agents: 1
  max_turns: 1
  max_total_turns: 1
  max_total_tokens: 600000
  max_retries: 1
  stage_contracts: off
  auto_commit_on_done: false
  auto_merge_on_done: false
  auto_merge_push_target: false
  after_done_failure_policy: block
codex:
  command: >-
    codex app-server --strict-config -c mcp_servers={{}} -c 'shell_environment_policy.inherit="all"'
    --enable skip_host_skill_discovery
    --disable apps --disable plugins --disable remote_plugin
    --disable browser_use --disable browser_use_external --disable browser_use_full_cdp_access
    --disable computer_use --disable image_generation --disable in_app_browser
    --disable in_app_local_automation --disable skill_search
    --disable skill_mcp_dependency_install --disable tool_call_mcp_elicitation
    --disable standalone_web_search --disable web_search_request --disable web_search_cached
    --disable memories --disable multi_agent --disable multi_agent_v2 --disable hooks
    --disable view_image
  model: gpt-6.1-sol
  reasoning_effort: medium
  approval_policy: never
  thread_sandbox: danger-full-access
  turn_sandbox_policy: danger-full-access
  turn_timeout_ms: 240000
pm:
  authoring: true
  exchange: true
  assistance: true
  handoff: true
  openspec_executable: {openspec}
  planning_workspace_root: {planning}
  operation_timeout_ms: 240000
  verification_procedures:
    focused: [/usr/bin/true]
  manual_verification_procedures:
    manual-ux:
      label: Installed manual review
      instructions: Review the immutable installed-fixture report.
      report: required
      artifacts: optional
---
This is a disposable installed-package execution proof for {{{{ issue.identifier }}}}.
Implement exactly one bounded change in the current worker worktree:
1. Create `implemented.txt` containing exactly `verified` followed by a newline.
2. Run `/usr/bin/true` as the task check.
3. Create `.symphony-artifacts/proof.txt` containing exactly `worker artifact` followed by a newline.
4. Commit `implemented.txt` with message `{{{{ issue.identifier }}}}: installed worker proof`.
5. Move the host ticket to Done with:
   `{sys.executable} -m symphony board mv {{{{ issue.identifier }}}} Done --workflow {workflow}`
Do not modify any other product file, push, merge, deploy, or archive. Finish only after
the commit exists and the board command succeeds.
""")
    if phase_reach:
        # Two active lanes force one genuine phase transition: the Review
        # lane is served by a freshly built backend whose process launch
        # must pass the PM permit guard again.
        source = workflow.read_text(encoding="utf-8")
        source = source.replace("  active_states: [Todo]\n", "  active_states: [Todo, Review]\n", 1)
        source = source.replace("  max_turns: 1\n  max_total_turns: 1\n", "  max_turns: 2\n  max_total_turns: 2\n", 1)
        source = source.replace(
            "5. Move the host ticket to Done with:\n"
            "   `{sys.executable} -m symphony board mv {{{{ issue.identifier }}}} Done --workflow {workflow}`\n".format(sys=sys, workflow=workflow),
            "Current lane: {{{{ issue.state }}}}.\n"
            "5. If the current lane is Todo, finish steps 1-4, then move the host ticket to Review with\n"
            "   `{sys.executable} -m symphony board mv {{{{ issue.identifier }}}} Review --workflow {workflow}`\n"
            "   and end your turn. If the current lane is Review, do not repeat steps 1-4: confirm\n"
            "   `implemented.txt` is committed, then move the host ticket to Done with\n"
            "   `{sys.executable} -m symphony board mv {{{{ issue.identifier }}}} Done --workflow {workflow}`\n".format(sys=sys, workflow=workflow),
            1,
        )
        if "active_states: [Todo, Review]" not in source or "max_turns: 2" not in source or "Current lane:" not in source:
            raise AssertionError("phase-reach workflow rewrite did not apply")
        workflow.write_text(source, encoding="utf-8")
    run("git", "add", ".", cwd=project)
    write(project / "reports" / "manual-ux.json", '{"result":"passed","review":"installed"}\n')
    run("git", "add", ".", cwd=project)
    run("git", "commit", "-qm", "fixture base", cwd=project)
    run("git", "worktree", "add", "-q", str(planning), "-b", "pm-planning-e2e", cwd=project)
    peer = root / "same-slug-peer"
    peer.mkdir()
    run("git", "init", "-q", cwd=peer)
    (peer / "kanban").mkdir()
    peer_workflow = peer / "WORKFLOW.md"
    write(peer_workflow, """---
tracker:
  kind: file
  board_root: ./kanban
  project_slug: pm-installed-e2e
  active_states: [Todo]
  terminal_states: [Done, Cancelled]
agent:
  kind: codex
---
Same tracker slug, distinct registered host identity.
""")
    peer_port = port - 1 if port > 1 else port + 1
    projects = root / "projects.json"
    projects.write_text(json.dumps({"version": 1, "projects": [
        {"id": "registered-pm-installed-e2e", "name": "PM installed E2E", "git_repo": str(project), "workflow": str(workflow), "host": "127.0.0.1", "port": port},
        {"id": "registered-pm-installed-e2e-peer", "name": "PM same-slug peer", "git_repo": str(peer), "workflow": str(peer_workflow), "host": "127.0.0.1", "port": peer_port},
    ]}, indent=2), encoding="utf-8")
    server = root / "run_server.py"
    write(server, """from aiohttp import web
import asyncio
import os
from pathlib import Path
from symphony.orchestrator import Orchestrator
from symphony.orchestrator.run_registry import RunRegistry
from symphony.pm.routes import PM_DEVELOPMENT_VERIFICATION_KEY
from symphony.server import build_app
from symphony.workflow import WorkflowState
import sys
state=WorkflowState(Path(sys.argv[1])); cfg,error=state.reload(); assert cfg is not None and error is None,error
handoff_fault=Path(sys.argv[1]).parent / ".inject-handoff-staging-crash"
original_create_handoff=RunRegistry.create_pm_handoff_batch
def create_handoff_then_crash(self,*args,**kwargs):
    result=original_create_handoff(self,*args,**kwargs)
    if handoff_fault.exists():
        handoff_fault.unlink()
        os._exit(86)
    return result
RunRegistry.create_pm_handoff_batch=create_handoff_then_crash
import json, time
from symphony.pm.authority import PmAuthority
spawn_log=Path(sys.argv[1]).parent / ".pm-guarded-spawns.jsonl"
original_launch_guard=PmAuthority.launch_guard
def recorded_launch_guard(self, permit, issue_loader, run_id, workspace_path=None):
    guard=original_launch_guard(self, permit, issue_loader, run_id, workspace_path)
    from contextlib import asynccontextmanager
    @asynccontextmanager
    async def wrapped():
        before=self.registry.get_pm_dispatch_permit(permit.permit_id)
        async with guard() as publish_started:
            async def publish(pid):
                await publish_started(pid)
                after=self.registry.get_pm_dispatch_permit(permit.permit_id)
                with spawn_log.open("a", encoding="utf-8") as stream:
                    stream.write(json.dumps({"ticket": permit.ticket_identifier, "run_id": run_id, "permit_id": permit.permit_id, "state_before": before.state if before else None, "state_after": after.state if after else None, "pid": pid, "at": time.time()}) + "\\n")
            yield publish
    return wrapped
PmAuthority.launch_guard=recorded_launch_guard
orchestrator=Orchestrator(state)
app=build_app(orchestrator)
app[PM_DEVELOPMENT_VERIFICATION_KEY]=True
dispatch_marker=Path(sys.argv[1]).parent / ".enable-dispatch"
async def enable_dispatch():
    while not dispatch_marker.exists(): await asyncio.sleep(0.05)
    await orchestrator.start()
async def start(_app): _app["dispatch_task"] = asyncio.create_task(enable_dispatch())
async def stop(_app):
    task = _app.get("dispatch_task")
    if task is not None and not task.done():
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
    await orchestrator.stop()
app.on_startup.append(start)
app.on_cleanup.append(stop)
web.run_app(app,host='127.0.0.1',port=int(sys.argv[2]),print=None)
""")
    return project, planning, server


def patch_document(api: Api, project: str, change: str, revision: str, document: str, source: str) -> str:
    body = envelope(project, change, revision, {"operations": [{"op": "replace_source", "source": source}]}, uuid.uuid4().hex)
    return api.call("PATCH", f"/changes/{change}/documents/{urllib.parse.quote(document, safe='')}", body)["revision"]


def semantic_id(source: str, kind: str) -> str:
    match = re.search(rf"<!--\s*symphony:{kind}=([A-Z]+-[0-9A-F]{{8}})\s*-->", source)
    if match is None:
        raise AssertionError(f"server-minted {kind} identity is missing")
    return match.group(1)


def await_worker_run(project_root: Path, ticket: str, timeout_seconds: float = 300) -> tuple[str, str, Path]:
    registry = RunRegistry(project_root / ".symphony" / "state.db")
    deadline = time.monotonic() + timeout_seconds
    last: str = "no run observed"
    try:
        while time.monotonic() < deadline:
            runs = registry.recent_runs(ticket, limit=10)
            if runs:
                record = runs[0]
                last = f"{record.status}: {record.failure_message or ''}".strip()
                worker_succeeded = (
                    record.status in {"completed", "normal"}
                    and record.state == "Done"
                    and record.commit_sha is not None
                )
                if worker_succeeded:
                    workspace = record.workspace_path.resolve()
                    implemented = workspace / "implemented.txt"
                    if implemented.read_text(encoding="utf-8") != "verified\n":
                        raise AssertionError("worker commit did not contain the approved implementation")
                    if run("git", "rev-parse", "HEAD", cwd=workspace) != record.commit_sha:
                        raise AssertionError("worker run commit does not match its preserved worktree")
                    return record.run_id, record.commit_sha, workspace
                if record.status not in {"active", "retrying"}:
                    raise AssertionError(f"Symphony worker ended without a verified commit: {last}")
            time.sleep(0.25)
    finally:
        registry.close()
    raise AssertionError(f"Symphony worker did not complete within {timeout_seconds}s: {last}")


def await_worker_release(
    project_root: Path,
    project_id: str,
    ticket: str,
    run_id: str,
    tested_commit: str,
    timeout_seconds: float = 60,
) -> dict[str, object]:
    registry = RunRegistry(project_root / ".symphony" / "state.db")
    deadline = time.monotonic() + timeout_seconds
    last = "run not found"
    try:
        while time.monotonic() < deadline:
            record = next(
                (
                    value
                    for value in registry.recent_runs(ticket, limit=10)
                    if value.run_id == run_id
                ),
                None,
            )
            if record is None:
                time.sleep(0.1)
                continue
            if record.status not in {"completed", "normal"}:
                last = f"run status={record.status}"
                time.sleep(0.1)
                continue
            assert record.commit_sha == tested_commit
            lease_active = registry.has_active_lease(record.issue_id)
            permit_active = registry.has_active_pm_dispatch(project_id, ticket)
            current_identity = (
                process_identity(record.backend_agent_pid)
                if record.backend_agent_pid is not None
                else None
            )
            group_exists = (
                process_group_exists(record.backend_agent_pid)
                if record.backend_agent_pid is not None
                else False
            )
            cleanup_confirmed = (
                record.backend_agent_pid is None
                or group_exists is False
                or (
                    record.backend_process_identity is not None
                    and current_identity is not None
                    and current_identity != record.backend_process_identity
                )
            )
            last = (
                f"lease={lease_active} permit={permit_active} "
                f"pid={record.backend_agent_pid} "
                f"cleanup_confirmed={cleanup_confirmed}"
            )
            if not lease_active and not permit_active and cleanup_confirmed:
                return {
                    "run_id": record.run_id,
                    "tested_commit": record.commit_sha,
                    "status": record.status,
                    "lease_active": False,
                    "pm_dispatch_active": False,
                    "recorded_backend_pid": record.backend_agent_pid,
                    "recorded_process_identity": record.backend_process_identity,
                    "current_process_identity": current_identity,
                    "process_group_exists": group_exists,
                    "owned_process_cleanup_confirmed": True,
                }
            time.sleep(0.1)
    finally:
        registry.close()
    raise AssertionError(
        "native worker authority/process cleanup did not finish within "
        f"{timeout_seconds}s: {last}"
    )


WBS_TREE_PRESS_BUDGET_SECONDS = 0.5
WBS_TREE_FILTER = "new "
_WBS_TREE_STATE = """([groupId, childId]) => {
  const rows = [...document.querySelectorAll('.pm-tabulator .tabulator-row')];
  const idOf = (row) => row.querySelector('.tabulator-cell[tabulator-field="task_id"]')?.textContent.trim();
  const group = rows.find((row) => idOf(row) === groupId);
  const control = group?.querySelector('.tabulator-data-tree-control');
  const active = document.activeElement;
  return {
    rows: rows.length,
    child_shown: rows.some((row) => idOf(row) === childId),
    role: control?.getAttribute('role') ?? null,
    label: control?.getAttribute('aria-label') ?? null,
    expanded: control?.getAttribute('aria-expanded') ?? null,
    focus: active === control ? 'control' : active === group ? 'group-row'
      : active?.closest?.('.tabulator-row') ? `row:${idOf(active.closest('.tabulator-row'))}` : active?.tagName,
    drawers: document.querySelectorAll('.pm-task-drawer').length,
    drawer_names: [...document.querySelectorAll('.pm-task-drawer')].map((item) => item.getAttribute('aria-label')),
  };
}"""


def _wbs_row(page, task_id: str):
    return page.locator(".pm-tabulator .tabulator-row").filter(
        has=page.locator(
            '.tabulator-cell[tabulator-field="task_id"]',
            has_text=re.compile(rf"^{re.escape(task_id)}$"),
        )
    )


def _wbs_command(page, name: str) -> str:
    with page.expect_response(
        lambda response: response.url.endswith("/wbs")
        and response.request.method == "PATCH"
    ) as info:
        button = page.get_by_role("button", name=name, exact=True)
        button.focus()
        page.keyboard.press("Enter")
    assert info.value.status == 200, info.value.text()
    return info.value.json()["impacted_ids"][0] if name.startswith("Add") else ""


def _wbs_open_and_close_drawer(page, task_id: str) -> None:
    row = _wbs_row(page, task_id)
    row.focus()
    page.keyboard.press("Enter")
    page.get_by_role("dialog", name=f"Task {task_id}").wait_for()
    page.keyboard.press("Escape")
    page.get_by_role("dialog", name=f"Task {task_id}").wait_for(state="hidden")


def prove_wbs_tree_keyboard(
    page, group_id: str, child_id: str, label: str = "Collapse or expand"
) -> dict[str, object]:
    """Assert WBS tree keyboard handling on a rendered parent/child pair.

    Repeated Enter/Space on the tree control must alternate collapse/expand
    with stable row counts, stay fast, keep focus on the rebuilt control, and
    never open a drawer; Enter on a row opens exactly one drawer even after
    pointer toggles (R32-01: listeners used to accumulate per re-format).
    """

    def state() -> dict:
        return page.evaluate(_WBS_TREE_STATE, [group_id, child_id])

    def press(key: str) -> tuple[dict, float]:
        started = time.monotonic()
        page.keyboard.press(key)
        page.evaluate("() => new Promise((resolve) => requestAnimationFrame(() => resolve()))")
        return state(), time.monotonic() - started

    def check(observed: dict, *, expanded: bool, focus: str, step: str) -> None:
        expected = {
            "rows": expanded_rows if expanded else collapsed_rows,
            "child_shown": expanded, "role": "button", "label": label,
            "expanded": str(expanded).lower(), "focus": focus, "drawers": 0,
        }
        actual = {key: observed[key] for key in expected}
        if collapsed_rows is None and not expanded:
            actual.pop("rows")
            expected.pop("rows")
        assert actual == expected, f"WBS tree {step}: expected {expected}, got {observed}"

    _wbs_row(page, child_id).wait_for()
    initial = state()
    expanded_rows, collapsed_rows = initial["rows"], None
    check(initial, expanded=True, focus=initial["focus"], step="initial render")
    _wbs_row(page, group_id).locator(".tabulator-data-tree-control").focus()
    presses: list[dict[str, object]] = []
    expanded = True
    for key in ("Enter", "Enter", "Space", "Space", "Enter", "Space", "Space", "Enter"):
        observed, seconds = press(key)
        expanded = not expanded
        check(observed, expanded=expanded, focus="control", step=f"control {key} #{len(presses) + 1}")
        if not expanded:
            collapsed_rows = collapsed_rows if collapsed_rows is not None else observed["rows"]
        assert collapsed_rows is not None and collapsed_rows < expanded_rows
        assert seconds < WBS_TREE_PRESS_BUDGET_SECONDS, f"WBS tree {key} press took {seconds:.3f}s"
        presses.append({"key": key, "expanded": expanded, "rows": observed["rows"], "seconds": round(seconds, 3)})
    _wbs_row(page, group_id).focus()
    for key, expanded in (("ArrowLeft", False), ("ArrowLeft", False), ("ArrowRight", True), ("ArrowRight", True)):
        observed, seconds = press(key)
        check(observed, expanded=expanded, focus="group-row", step=f"parent row {key}")
        assert seconds < WBS_TREE_PRESS_BUDGET_SECONDS, f"WBS tree {key} press took {seconds:.3f}s"
        presses.append({"key": key, "expanded": expanded, "rows": observed["rows"], "seconds": round(seconds, 3)})
    observed, _ = press("ArrowDown")
    assert observed["focus"] == f"row:{child_id}", observed
    observed, _ = press("ArrowUp")
    assert observed["focus"] == "group-row", observed
    for _ in range(4):
        _wbs_row(page, group_id).locator(".tabulator-data-tree-control").click()
    assert state()["child_shown"] is True
    row_enter: dict[str, object] = {}
    for task_id in (group_id, child_id):
        _wbs_row(page, task_id).focus()
        observed, _ = press("Enter")
        assert observed["drawers"] == 1 and observed["drawer_names"] == [f"Task {task_id}"], observed
        page.keyboard.press("Escape")
        page.get_by_role("dialog", name=f"Task {task_id}").wait_for(state="hidden")
        row_enter[task_id] = observed["drawers"]
    return {
        "group_id": group_id, "child_id": child_id,
        "expanded_rows": expanded_rows, "collapsed_rows": collapsed_rows,
        "control": {key: initial[key] for key in ("role", "label", "expanded")},
        "presses": presses, "row_enter_drawers": row_enter,
        "press_budget_seconds": WBS_TREE_PRESS_BUDGET_SECONDS,
    }


def exercise_wbs_tree_keyboard(page) -> dict[str, object]:
    """Create a disposable group/child in the Plan view, prove tree keys, remove both."""
    page.get_by_role("button", name="Plan", exact=True).click()
    group_id = _wbs_command(page, "Add group")
    search = page.get_by_label("Filter tasks")
    search.fill(WBS_TREE_FILTER)
    _wbs_open_and_close_drawer(page, group_id)
    child_id = _wbs_command(page, "Add child")
    proof = prove_wbs_tree_keyboard(page, group_id, child_id)
    for task_id in (child_id, group_id):
        _wbs_open_and_close_drawer(page, task_id)
        page.once("dialog", lambda dialog: dialog.accept())
        _wbs_command(page, "Delete")
        _wbs_row(page, task_id).wait_for(state="detached")
    search.fill("")
    return proof


class BrowserProof:
    """Drive the installed static UI while the disposable server is live."""

    def __init__(self, url: str, evidence_root: Path) -> None:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise RuntimeError("--browser requires the installed browser extra") from exc
        self.evidence_root = evidence_root
        self.evidence_root.mkdir(parents=True, exist_ok=True)
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(headless=True)
        self.page = self.browser.new_page(viewport={"width": 1366, "height": 768})
        # Secondary PM controls sit in collapsed "pm-more" disclosures. This
        # journey proves the actions, not the disclosures, so open each one as
        # it renders; the page remembers open disclosures across re-renders.
        self.page.add_init_script(
            "new MutationObserver(() => document.querySelectorAll('details.pm-more:not([open])')"
            ".forEach((details) => { details.open = true; }))"
            ".observe(document, { childList: true, subtree: true });"
        )
        self.page.goto(url, wait_until="networkidle")
        self.page.get_by_role("heading", name="Artifacts", exact=True).wait_for()
        self.actions: list[str] = ["opened_registered_project"]
        self.task_id = ""
        self.requirements_delta: dict[str, object] = {}
        self.wbs_tree_keyboard: dict[str, object] = {}

    def review_and_edit(self) -> None:
        page = self.page
        page.get_by_role("button", name="Idea", exact=True).click()
        page.get_by_role("button", name="Edit source", exact=True).click()
        editor = page.get_by_label("proposal source")
        editor.fill(editor.input_value().rstrip() + "\n\nUI review: installed browser journey.\n")
        previous_revision = page.locator(".pm-revision-bar code").inner_text()
        with page.expect_response(
            lambda response: "/documents/proposal" in response.url
            and response.request.method == "PATCH"
        ) as response_info:
            page.get_by_role("button", name="Save draft", exact=True).click()
        response = response_info.value
        if response.status != 200:
            raise AssertionError(
                f"browser proposal save failed: HTTP {response.status}: {response.text()}"
            )
        page.wait_for_function(
            "old => document.querySelector('.pm-revision-bar code')?.textContent !== old",
            arg=previous_revision,
        )
        self.actions.append("edited_and_saved_idea")
        page.get_by_role("button", name="Requirements", exact=True).click()
        try:
            page.locator(".pm-requirement-delta").wait_for(timeout=10_000)
        except Exception:
            page.screenshot(
                path=str(self.evidence_root / "requirements-failure.png"),
                full_page=True,
            )
            write(
                self.evidence_root / "requirements-failure.txt",
                page.locator("body").inner_text(),
            )
            raise
        delta = page.locator(".pm-requirement-delta")
        baseline_text = delta.locator(".pm-baseline").inner_text()
        proposed_text = delta.locator(".pm-document-preview").first.inner_text()
        assert "new capability" not in baseline_text.lower()
        assert "Existing accepted baseline" in baseline_text
        assert "Existing accepted baseline" in proposed_text
        assert "retain accepted behavior and explicit review boundaries" in proposed_text
        self.requirements_delta = {
            "accepted_baseline_exists": True,
            "same_requirement_modified": True,
            "baseline_excerpt": baseline_text[:500],
            "proposed_excerpt": proposed_text[:500],
        }
        self.actions.append("reviewed_requirements_delta")
        page.get_by_role("button", name="Design", exact=True).click()
        for name in ("Storyboard", "Architecture", "Integration"):
            page.get_by_role("tab", name=name, exact=True).click()
        self.actions.append("reviewed_storyboard_architecture_integration")
        page.get_by_role("button", name="Plan", exact=True).click()
        page.get_by_text(self.task_id, exact=True).first.wait_for()
        self.actions.append("reviewed_wbs")
        page.screenshot(path=str(self.evidence_root / "review.png"), full_page=True)

    def exercise_work_item_identity(self) -> None:
        page = self.page
        page.get_by_role("button", name="Plan", exact=True).click()
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ) as group_info:
            page.get_by_role("button", name="Add group", exact=True).click()
        group_payload = group_info.value.json()
        assert group_info.value.status == 200, group_info.value.text()
        group_id = group_payload["impacted_ids"][0]
        assert group_id.startswith("G-")
        page.get_by_text(group_id, exact=True).first.wait_for()
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ) as task_info:
            page.get_by_role("button", name="Add root", exact=True).click()
        task_payload = task_info.value.json()
        assert task_info.value.status == 200, task_info.value.text()
        task_id = task_payload["impacted_ids"][0]
        assert task_id.startswith("T-")
        page.get_by_text(task_id, exact=True).first.wait_for()
        page.get_by_text("New task", exact=True).first.click()
        drawer = page.get_by_role("dialog", name=f"Task {task_id}")
        drawer.wait_for()
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ):
            title = drawer.get_by_label("Title")
            title.fill("Browser minted task")
            title.press("Tab")
        page.get_by_role("button", name="Close", exact=True).click()
        page.get_by_text("Browser minted task", exact=True).first.click()
        page.get_by_role("button", name="Close", exact=True).click()
        previous_revision = page.locator(".pm-revision-bar code").inner_text()
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ):
            page.get_by_role("button", name="Move up", exact=True).click()
        page.wait_for_function(
            "old => document.querySelector('.pm-revision-bar code')?.textContent !== old",
            arg=previous_revision,
        )
        page.once("dialog", lambda dialog: dialog.accept())
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ):
            page.get_by_role("button", name="Delete", exact=True).click()
        page.get_by_text("Browser minted task", exact=True).first.wait_for(state="hidden")
        page.wait_for_function(
            "() => [...document.querySelectorAll('button')].some((item) => item.textContent.trim() === 'Undo' && !item.disabled)"
        )
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ) as undo_info:
            page.get_by_role("button", name="Undo", exact=True).click()
        assert undo_info.value.status == 200, undo_info.value.text()
        page.get_by_text("Browser minted task", exact=True).first.wait_for()
        page.wait_for_function(
            "() => [...document.querySelectorAll('button')].some((item) => item.textContent.trim() === 'Redo' && !item.disabled)"
        )
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ) as redo_info:
            page.get_by_role("button", name="Redo", exact=True).click()
        assert redo_info.value.status == 200, redo_info.value.text()
        redo_revision = redo_info.value.json()["revision"]
        page.wait_for_function(
            "revision => document.querySelector('.pm-revision-bar code')?.textContent === revision",
            arg=redo_revision,
        )
        page.get_by_text("Browser minted task", exact=True).first.wait_for(state="hidden")
        page.get_by_text(group_id, exact=True).first.click()
        group_drawer = page.get_by_role("dialog", name=f"Task {group_id}")
        group_drawer.wait_for()
        group_drawer.get_by_role("button", name="Close", exact=True).click()
        group_drawer.wait_for(state="hidden")
        delete_button = page.get_by_role("button", name="Delete", exact=True)
        assert delete_button.is_enabled()
        page.once("dialog", lambda dialog: dialog.accept())
        with page.expect_response(
            lambda response: response.url.endswith("/wbs")
            and response.request.method == "PATCH"
        ) as group_delete_info:
            delete_button.click()
        assert group_delete_info.value.status == 200, group_delete_info.value.text()
        page.get_by_text(group_id, exact=True).first.wait_for(state="hidden")
        self.actions.append("created_renamed_reordered_deleted_and_restored_server_work_ids")

    def exercise_wbs_tree_keyboard(self) -> None:
        self.wbs_tree_keyboard = exercise_wbs_tree_keyboard(self.page)
        self.actions.append("toggled_wbs_tree_by_keyboard")

    def export_workbook(self) -> Path:
        with self.page.expect_download() as download_info:
            self.page.get_by_role("button", name="Download export", exact=True).click()
        target = self.evidence_root / "planning-base.xlsx"
        download_info.value.save_as(target)
        from openpyxl import load_workbook

        workbook = load_workbook(target)
        changed = False
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows())
            if not rows:
                continue
            headers = {str(cell.value or "").strip(): index for index, cell in enumerate(rows[0], 1)}
            task_column = headers.get("Task ID")
            title_column = headers.get("Title")
            if task_column is None or title_column is None:
                continue
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, task_column).value == self.task_id:
                    sheet.cell(row, title_column).value = "Workbook reviewed title"
                    changed = True
                    break
        if not changed:
            raise AssertionError("exported workbook did not expose the stable task/title fields")
        reviewed = self.evidence_root / "planning-reviewed.xlsx"
        workbook.save(reviewed)
        self.actions.append("exported_and_edited_xlsx")
        return reviewed

    def import_conflict(self, workbook: Path) -> None:
        page = self.page
        page.reload(wait_until="networkidle")
        page.get_by_role("button", name="Plan", exact=True).click()
        page.get_by_label("Choose XLSX workbook").set_input_files(str(workbook))
        with page.expect_response(
            lambda response: response.url.endswith("/imports/preview")
            and response.request.method == "POST"
        ) as preview_info:
            page.get_by_role("button", name="Preview import", exact=True).click()
        assert preview_info.value.status == 200, preview_info.value.text()
        page.locator(".pm-import-preview").wait_for()
        page.get_by_label(f"Resolve {self.task_id}:title").select_option("imported")
        with page.expect_response(
            lambda response: "/imports/" in response.url
            and response.url.endswith("/apply")
            and response.request.method == "POST"
        ) as apply_info:
            page.get_by_role("button", name="Apply reviewed import", exact=True).click()
        assert apply_info.value.status == 200, apply_info.value.text()
        self.actions.append("resolved_and_applied_xlsx_conflict")

    def create_change(self, workspace: Path) -> None:
        answers = iter(("installed-proof", "Installed proof", str(workspace)))

        def answer(dialog) -> None:
            dialog.accept(next(answers))

        self.page.on("dialog", answer)
        self.page.get_by_role("button", name="New change", exact=True).click()
        self.page.locator(".pm-change-select option").filter(
            has_text="installed-proof"
        ).wait_for(state="attached")
        self.page.remove_listener("dialog", answer)
        self.actions.append("created_change_from_idea_in_ui")

    def approve_revision(self) -> None:
        page = self.page
        page.get_by_role("button", name="Validate", exact=True).click()
        page.get_by_role("dialog", name="Validation results").wait_for()
        page.get_by_role("button", name="Close", exact=True).click()
        approval_bar = page.locator(".pm-revision-bar").first.element_handle()
        assert approval_bar is not None
        page.once("dialog", lambda dialog: dialog.accept())
        with page.expect_response(
            lambda response: response.url.endswith("/approvals")
            and response.request.method == "POST"
        ) as approval_info:
            page.get_by_role("button", name="Approve revision", exact=True).click()
        approval_response = approval_info.value
        assert approval_response.status == 201, approval_response.text()
        approved_revision = approval_response.json()["revision"]
        page.wait_for_function("element => !element.isConnected", arg=approval_bar)
        refreshed_revision = page.locator(".pm-revision-bar code").first
        refreshed_revision.wait_for()
        assert refreshed_revision.get_attribute("title") == approved_revision
        self.actions.extend(("validated_in_ui", "approved_exact_revision_in_ui"))

    def _open_handoff_confirmation(self) -> None:
        page = self.page
        page.get_by_role("button", name="Plan", exact=True).click()
        with page.expect_response(
            lambda response: response.url.endswith("/handoffs/preview")
            and response.request.method == "POST"
        ) as preview_info:
            page.get_by_role("button", name="Preview approved handoff", exact=True).click()
        preview_response = preview_info.value
        assert preview_response.status == 200, preview_response.text()
        page.get_by_role("dialog", name="Confirm native handoff").wait_for()

    def _set_operation_id(self, operation_id: str) -> None:
        value = str(uuid.UUID(hex=operation_id))
        self.page.evaluate(
            """value => {
              crypto.__symphonyOriginalRandomUUID ||= crypto.randomUUID.bind(crypto);
              crypto.randomUUID = () => value;
            }""",
            value,
        )

    def _restore_operation_ids(self) -> None:
        self.page.evaluate(
            """() => {
              if (crypto.__symphonyOriginalRandomUUID) {
                crypto.randomUUID = crypto.__symphonyOriginalRandomUUID;
                delete crypto.__symphonyOriginalRandomUUID;
              }
            }"""
        )

    def start_handoff_with_lost_response(self, operation_id: str) -> None:
        page = self.page
        self._set_operation_id(operation_id)
        self._open_handoff_confirmation()
        try:
            with page.expect_response(
                lambda response: response.url.endswith("/handoffs")
                and response.request.method == "POST",
                timeout=10_000,
            ):
                page.get_by_role(
                    "button", name="Create native tickets", exact=True
                ).click()
        except Exception:
            self.actions.append("lost_staging_handoff_response")
            return
        raise AssertionError("handoff response was not interrupted after STAGING")

    def recover_handoff(self, url: str, operation_id: str) -> dict[str, object]:
        page = self.page
        page.goto(url, wait_until="networkidle")
        page.get_by_role("heading", name="Artifacts", exact=True).wait_for()
        self._set_operation_id(operation_id)
        self._open_handoff_confirmation()
        handoff_bar = page.locator(".pm-revision-bar").first.element_handle()
        assert handoff_bar is not None
        with page.expect_response(
            lambda response: response.url.endswith("/handoffs")
            and response.request.method == "POST"
        ) as publish_info:
            page.get_by_role(
                "button", name="Create native tickets", exact=True
            ).click()
        assert publish_info.value.status == 201, publish_info.value.text()
        page.wait_for_function("element => !element.isConnected", arg=handoff_bar)
        page.locator(".pm-revision-bar").first.wait_for()
        self._restore_operation_ids()
        self.actions.append("confirmed_native_handoff_in_ui")
        return publish_info.value.json()

    def accept_verified(self) -> None:
        page = self.page
        page.get_by_role("button", name="Code Files", exact=True).click()
        card = page.locator(".pm-ticket-card").filter(has_text=self.task_id)
        card.get_by_text(
            "Verification evidence · 1 receipts · Host executed", exact=True
        ).click()
        with page.expect_download() as download_info:
            card.get_by_role(
                "button", name="focused · proof.txt", exact=True
            ).click()
        artifact = self.evidence_root / "worker-artifact.txt"
        download_info.value.save_as(artifact)
        assert artifact.read_text(encoding="utf-8") == "worker artifact\n"
        self.actions.append("opened_real_worker_artifact_in_ui")
        card.get_by_role(
            "button", name="Review verification", exact=True
        ).click()
        page.get_by_role("dialog", name="Acceptance review").wait_for()
        acceptance_bar = page.locator(".pm-revision-bar").first.element_handle()
        assert acceptance_bar is not None
        with page.expect_response(
            lambda response: response.url.endswith("/acceptance")
            and response.request.method == "POST"
        ) as acceptance_info:
            page.get_by_role(
                "button", name="Accept verified revision", exact=True
            ).click()
        assert acceptance_info.value.status == 201, acceptance_info.value.text()
        page.wait_for_function("element => !element.isConnected", arg=acceptance_bar)
        page.locator(".pm-revision-bar").first.wait_for()
        self.actions.extend(("reviewed_integrated_evidence_in_ui", "accepted_verified_revision_in_ui"))

    def archive(self) -> None:
        page = self.page
        page.get_by_role("button", name="Plan", exact=True).click()
        page.get_by_role("button", name="Preview archive", exact=True).click()
        dialog = page.get_by_role("dialog", name="Confirm archive")
        dialog.wait_for()
        page.screenshot(path=str(self.evidence_root / "archive-preview.png"), full_page=True)
        page.get_by_role("button", name="Archive reviewed revision", exact=True).click()
        dialog.wait_for(state="hidden", timeout=30_000)
        archived_option = page.locator(
            ".pm-change-select option", has_text="installed-proof · Archived"
        )
        assert archived_option.count() == 1
        assert archived_option.first.evaluate("option => option.selected") is True
        assert page.locator(".banner-danger").count() == 0
        self.actions.append("confirmed_archive_in_ui")

    def review_archived_history(self, url: str, task_id: str) -> None:
        page = self.page
        page.goto(url, wait_until="networkidle")
        page.get_by_role("heading", name="Artifacts", exact=True).wait_for()
        page.get_by_text("Archived", exact=True).first.wait_for()
        page.get_by_role("button", name="Idea", exact=True).click()
        page.get_by_role("heading", name="Read-only document", exact=True).wait_for()
        assert page.get_by_role("button", name="Edit source", exact=True).count() == 0
        page.get_by_role("button", name="Plan", exact=True).click()
        page.get_by_text(task_id, exact=True).first.wait_for()
        page.get_by_role("button", name="Code Files", exact=True).click()
        card = page.locator(".pm-ticket-card").filter(has_text=task_id)
        evidence = card.get_by_text(
            "Verification evidence · 1 receipts · Host executed", exact=True
        )
        evidence.wait_for()
        evidence.click()
        with page.expect_download() as download_info:
            card.get_by_role(
                "button", name="focused · proof.txt", exact=True
            ).click()
        artifact = self.evidence_root / "archived-worker-artifact.txt"
        download_info.value.save_as(artifact)
        assert artifact.read_text(encoding="utf-8") == "worker artifact\n"
        page.screenshot(
            path=str(self.evidence_root / "archived-history.png"), full_page=True
        )
        self.actions.append("reopened_original_change_id_as_readonly_history")

    def accept_and_archive(self) -> None:
        self.accept_verified()
        self.archive()

    def close(self) -> dict[str, object]:
        self.browser.close()
        self.playwright.stop()
        return {
            "driver": "playwright",
            "actions": self.actions,
            "requirements_delta": self.requirements_delta,
            "wbs_tree_keyboard": self.wbs_tree_keyboard,
            "screenshots": [
                str(self.evidence_root / "review.png"),
                str(self.evidence_root / "archive-preview.png"),
            ],
        }


def journey(
    root: Path, openspec: Path, port: int, *, assistance: bool, browser: bool,
    phase_reach: bool = False,
) -> dict:
    project_root, planning, server_script = setup(root, openspec, port, phase_reach=phase_reach)
    doctor = subprocess.run(
        [sys.executable, "-m", "symphony", "doctor", "--no-color", str(project_root / "WORKFLOW.md")],
        cwd=project_root,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=60,
    )
    if doctor.returncode != 0:
        raise RuntimeError(f"installed workflow doctor failed:\n{doctor.stdout}")
    environment = {**os.environ, "SYMPHONY_PROJECTS_FILE": str(root / "projects.json")}
    server_log = (root / "server.log").open("w", encoding="utf-8")
    server = subprocess.Popen([sys.executable, str(server_script), str(project_root / "WORKFLOW.md"), str(port)], cwd=project_root, env=environment, stdout=server_log, stderr=subprocess.STDOUT, text=True)
    api = Api(port)
    browser_proof: BrowserProof | None = None
    try:
        wait_for_api(api, server, root / "server.log")

        project = "registered-pm-installed-e2e"
        peer_project = "registered-pm-installed-e2e-peer"
        capabilities = api.call("GET", "/capabilities")
        assert capabilities["project_id"] == project
        same_slug_stale_refused = False
        try:
            api.call("GET", f"/changes?project_id={peer_project}")
        except RuntimeError as exc:
            same_slug_stale_refused = (
                "HTTP 409" in str(exc) and "wrong_project" in str(exc)
            )
        assert same_slug_stale_refused
        client_change_id_rejected = False
        try:
            api.call("POST", "/changes", {
                "project_id": project, "change_id": uuid.uuid4().hex,
                "workspace_path": str(planning), "locator": "client-id-attack",
                "title": "Client supplied identity",
            })
        except RuntimeError as exc:
            client_change_id_rejected = "server_minted_id_required" in str(exc)
        assert client_change_id_rejected
        if browser:
            browser_proof = BrowserProof(
                f"http://127.0.0.1:{port}/#/artifacts",
                root / "browser-evidence",
            )
            browser_proof.create_change(planning)
            changes = api.call("GET", f"/changes?project_id={project}")["changes"]
            assert len(changes) == 1
            change = changes[0]["change_id"]
        else:
            created_change = api.call("POST", "/changes", {
                "project_id": project, "workspace_path": str(planning),
                "locator": "installed-proof", "title": "Installed proof",
            })
            change = created_change["change_id"]
        current = api.call("GET", f"/changes/{change}?project_id={project}")
        assert any(item["document_id"] == "spec:requirements" for item in current["documents"])
        assert (planning / "openspec" / "changes" / "installed-proof" / ".openspec.yaml").read_text().startswith("schema: symphony-pm\n")

        seeded = {
            document_id: api.call(
                "GET",
                f"/changes/{change}/documents/{urllib.parse.quote(document_id, safe='')}?project_id={project}",
            )["source"]
            for document_id in ("proposal", "ux", "design", "spec:requirements")
        }
        fresh_requirement_id = semantic_id(seeded["spec:requirements"], "req")
        fresh_scenario_id = semantic_id(seeded["spec:requirements"], "scenario")
        baseline_source = (
            planning / "openspec" / "specs" / "requirements" / "spec.md"
        ).read_text(encoding="utf-8")
        requirement_id = semantic_id(baseline_source, "req")
        scenario_id = semantic_id(baseline_source, "scenario")
        assert requirement_id != fresh_requirement_id
        assert scenario_id != fresh_scenario_id
        storyboard_id = semantic_id(seeded["ux"], "storyboard")
        contract_id = semantic_id(seeded["design"], "contract")
        decision_id = semantic_id(seeded["proposal"], "decision")
        sources = {
            "proposal": f"""# Installed proof

## Problem
Product teams need one trustworthy installed-package planning path.

## Intended users
Product owners and developers reviewing approved work.

## Desired outcome
A reviewable plan reaches explicit acceptance, local delivery, and archive.

## Scope
One disposable PM change and one native Symphony worker ticket.

## Non-goals
No push, production merge, deploy, or real-project archive.

## Constraints
Use the installed wheel, pinned OpenSpec, and manual delivery policy.

## Success criteria
The worker commit passes host verification and receives human acceptance.

## Assumptions
Local Codex authentication and Chromium are available.

## Open decisions

### Decision: Use local disposable delivery
<!-- symphony:decision={decision_id} -->
- Owner: Local operator
- Impact: Low
- Status: Accepted
- Reason: Keeps the proof reversible and isolated.

## Affected capabilities
- requirements
""",
            "ux": f"""# UX and storyboard

## Actors and journeys
The product owner reviews the plan before execution.

## Screen inventory
Artifacts Idea, Requirements, Design, Plan, and Code Files tabs.

## Loading, empty, error, and permission states
Loading and failures remain visible without granting authority.

## Responsive and accessibility behavior
Essential review and actions remain keyboard operable.

## Storyboard

### Storyboard: Installed review journey
<!-- symphony:storyboard={storyboard_id} -->
- Screen/route: Artifacts
- Actor/goal: Product owner approves exact scope
- Initial state: Validated draft
- UI/action/response: Review, approve, hand off, verify, and accept
- Success/alternate/failure: Archive only after accepted evidence
""",
            "design": f"""# Technical design

## Architecture and ownership
```mermaid
flowchart LR
 Browser --> API
 API --> Worktree
```

## Data model and migration
Use existing PM registry records without an additional service.

## API and integration contracts

### Contract: Native file-board handoff
<!-- symphony:contract={contract_id} -->
- Owner/boundary: Symphony PM host and file tracker
- Auth/schema/errors/version/idempotency: Existing guarded PM API
- Retry/partial failure: Durable operation and handoff records
- Verification: Host-run approved command at the worker commit

## Security, concurrency, and operations
Manual delivery, authority fencing, and disposable local paths.

## Verification and rollback
Verify the exact commit; discard the disposable fixture after proof.
""",
            "spec:requirements": f"""# Requirements

## MODIFIED Requirements

### Requirement: Existing accepted baseline
<!-- symphony:req={requirement_id} -->

The existing system SHALL retain accepted behavior and explicit review boundaries.

#### Scenario: Baseline remains available
<!-- symphony:scenario={scenario_id} -->
- **WHEN** a proposed change is opened
- **THEN** the accepted baseline remains visible for comparison

#### Scenario: Verified handoff
<!-- symphony:scenario={fresh_scenario_id} -->
- **GIVEN** an approved planning revision
- **WHEN** its native worker commit passes host verification
- **THEN** a human can explicitly accept and archive it
""",
        }
        for document, source in sources.items():
            current["revision"] = patch_document(api, project, change, current["revision"], document, source)
            current = api.call("GET", f"/changes/{change}?project_id={project}")

        task_ids: dict[str, str] = {}
        task_definitions = (
            ("worker", "Prove installed lifecycle", 10, [
                {"procedure_id": "focused", "kind": "native"},
            ]),
            ("native_no_run", "Host-executed no-run verification", 20, [
                {"procedure_id": "focused", "kind": "native"},
            ]),
            ("manual_only", "Manual-only completion evidence", 30, [
                {"procedure_id": "manual-ux", "kind": "manual"},
            ]),
            ("mixed", "Mixed no-run completion evidence", 40, [
                {"procedure_id": "focused", "kind": "native"},
                {"procedure_id": "manual-ux", "kind": "manual"},
            ]),
        )
        for key, title, order, policy in task_definitions:
            created_task = api.call(
                "PATCH", f"/changes/{change}/wbs",
                envelope(project, change, current["revision"], {"operations": [{
                    "op": "create_task",
                    "task": {
                        "kind": "Backend", "title": title,
                        "parent_id": None, "order": order, "dependencies": [],
                        "requirement_refs": [requirement_id, fresh_scenario_id],
                        "storyboard_refs": [storyboard_id],
                        "design_refs": [contract_id],
                        "description": "Create bounded evidence for the installed lifecycle.",
                        "acceptance": ["The exact tested commit is accepted."],
                        "verification": ["Use the exact approved procedure policy."],
                        "verification_policy": policy,
                        "integration_impact": "Native file tracker ticket and durable evidence receipt.",
                    },
                }]}, uuid.uuid4().hex),
            )
            task_ids[key] = created_task["impacted_ids"][0]
            current = api.call("GET", f"/changes/{change}?project_id={project}")
        task_id = task_ids["worker"]
        tasks_source = api.call(
            "GET", f"/changes/{change}/documents/tasks?project_id={project}"
        )["source"]
        raw_identity_rejected = False
        try:
            patch_document(
                api, project, change, current["revision"], "tasks",
                tasks_source.replace(task_id, "T-ATTACK01"),
            )
        except RuntimeError as exc:
            raw_identity_rejected = "work_item_id_replacement" in str(exc)
        assert raw_identity_rejected
        if browser_proof is not None:
            browser_proof.task_id = task_id

        if browser:
            assert browser_proof is not None
            browser_proof.page.goto(
                f"http://127.0.0.1:{port}/?pmTab=idea&pmChange={change}#/artifacts",
                wait_until="networkidle",
            )
            browser_proof.page.reload(wait_until="networkidle")
            browser_proof.page.get_by_role(
                "heading", name="Artifacts", exact=True
            ).wait_for()
            browser_proof.exercise_work_item_identity()
            browser_proof.exercise_wbs_tree_keyboard()
            browser_proof.review_and_edit()
            current = api.call("GET", f"/changes/{change}?project_id={project}")
            reviewed_workbook = browser_proof.export_workbook()
            api.call(
                "PATCH",
                f"/changes/{change}/wbs",
                envelope(
                    project,
                    change,
                    current["revision"],
                    {"operations": [{
                        "op": "set_field", "task_id": task_id,
                        "field": "title", "value": "Browser current title",
                    }]},
                    uuid.uuid4().hex,
                ),
            )
            browser_proof.import_conflict(reviewed_workbook)
            current = api.call("GET", f"/changes/{change}?project_id={project}")

        assistance_receipt: dict = {"requested": assistance}
        if assistance:
            operation = uuid.uuid4().hex
            proposal_body = {
                **envelope(project, change, current["revision"], operation=operation),
                "action": "draft_artifact",
                "document_id": "proposal",
                "impacted_document_ids": ["proposal"],
                "instruction": (
                    "Return exactly one unified-diff hunk for proposal.md. The hunk header must be "
                    "@@ -9,2 +9,3 @@ and its two context lines must be exactly "
                    "' ## Desired outcome' and ' A reviewable plan reaches explicit acceptance, "
                    "local delivery, and archive.'. Add '+Review note: the installed proof was "
                    "reviewed.' after those context lines. "
                    "Do not include file headers, prose, or tools."
                ),
            }
            proposal = api.call("POST", f"/changes/{change}/planning-actions", proposal_body)
            assert proposal["state"] == "completed" and proposal["result"]["write_or_dispatch_performed"] is False
            assert proposal["result"]["proposals"], proposal
            assert {
                item["document_id"] for item in proposal["result"]["proposals"]
            } == {"proposal"}, proposal
            rejected = False
            try:
                bad_apply = {
                    **envelope(project, change, current["revision"], operation=uuid.uuid4().hex),
                    "proposal_digest": "sha256:" + ("0" * 64),
                }
                api.call("POST", f"/changes/{change}/planning-actions/{operation}/apply", bad_apply)
            except RuntimeError as exc:
                rejected = "confirmation_mismatch" in str(exc)
            assert rejected, "the host accepted an unconfirmed planning proposal"
            apply_body = {**envelope(project, change, current["revision"], operation=uuid.uuid4().hex), "proposal_digest": proposal["result"]["proposal_digest"]}
            api.call("POST", f"/changes/{change}/planning-actions/{operation}/apply", apply_body)
            current = api.call("GET", f"/changes/{change}?project_id={project}")
            proposal_source = api.call(
                "GET", f"/changes/{change}/documents/proposal?project_id={project}"
            )["source"]
            assert "Review note: the installed proof was reviewed." in proposal_source

            cancel_operation = uuid.uuid4().hex
            cancel_result: dict[str, object] = {}

            def request_long_proposal() -> None:
                try:
                    cancel_result["response"] = api.call(
                        "POST",
                        f"/changes/{change}/planning-actions",
                        {
                            **envelope(
                                project,
                                change,
                                current["revision"],
                                operation=cancel_operation,
                            ),
                            "action": "find_gaps",
                            "document_id": "proposal",
                            "impacted_document_ids": ["proposal"],
                            "instruction": (
                                "Carefully review every sentence and consider many alternatives before "
                                "returning a proposal. Do not use tools or modify anything."
                            ),
                        },
                    )
                except Exception as exc:  # cancellation may close the in-flight request
                    cancel_result["request_error"] = str(exc)

            request_thread = threading.Thread(target=request_long_proposal, daemon=True)
            request_thread.start()
            for _ in range(200):
                try:
                    state = api.call("GET", f"/operations/{cancel_operation}")
                except RuntimeError:
                    time.sleep(0.05)
                    continue
                if state["state"] in {"pending", "running"}:
                    break
                time.sleep(0.05)
            else:
                raise AssertionError("planning cancellation operation never started")
            time.sleep(0.5)
            cancelled = api.call("POST", f"/operations/{cancel_operation}/cancel", {})
            request_thread.join(timeout=30)
            assert not request_thread.is_alive(), "cancelled planning request did not terminate"
            final_cancel = api.call("GET", f"/operations/{cancel_operation}")
            assert cancelled["state"] == "cancelled" and final_cancel["state"] == "cancelled"
            assistance_receipt.update({
                "operation_id": operation,
                "applied": True,
                "invalid_digest_rejected": True,
                "cancel_operation_id": cancel_operation,
                "cancelled": True,
                "model": "gpt-6.1-sol",
                "effort": "medium",
                "proposal_digest": proposal["result"]["proposal_digest"],
            })

        if browser_proof is not None:
            browser_proof.page.reload(wait_until="networkidle")
            browser_proof.page.get_by_role(
                "heading", name="Artifacts", exact=True
            ).wait_for()

        validation = api.call("POST", f"/changes/{change}/validate", envelope(project, change, current["revision"]))
        assert validation["valid"] is True, validation
        handoff_recovery: dict[str, object] = {"staging_response_lost": False}
        if browser_proof is not None:
            browser_proof.approve_revision()
            current = api.call("GET", f"/changes/{change}?project_id={project}")
            registry_probe = RunRegistry(project_root / ".symphony" / "state.db")
            approval_row = registry_probe.get_active_pm_approval(project, change)
            registry_probe.close()
            assert approval_row is not None
            approval = {"approval_id": approval_row["approval_id"]}

            handoff_operation = uuid.uuid4().hex
            (project_root / ".inject-handoff-staging-crash").write_text(
                "crash after durable staging\n", encoding="utf-8"
            )
            browser_proof.start_handoff_with_lost_response(handoff_operation)
            server.wait(timeout=15)
            assert server.returncode == 86

            staged_registry = RunRegistry(project_root / ".symphony" / "state.db")
            try:
                staged = staged_registry.get_pm_handoff_by_operation(
                    handoff_operation
                )
                assert staged is not None and staged["state"] == "STAGING"
                staged_items = staged_registry.get_pm_handoff_items(
                    str(staged["handoff_id"])
                )
                assert len(staged_items) == 4
                assert all(item["state"] == "STAGING" for item in staged_items)
                assert all(
                    not staged_registry.recent_runs(
                        str(item["deterministic_ticket_id"]), limit=10
                    )
                    for item in staged_items
                )
            finally:
                staged_registry.close()
            state = WorkflowState(project_root / "WORKFLOW.md")
            config, error = state.reload()
            assert config is not None and error is None, error
            staged_tracker = FileBoardTracker(config.tracker)
            assert all(
                staged_tracker.fetch_issue_full_by_id(
                    str(item["deterministic_ticket_id"])
                ) is None
                for item in staged_items
            )

            server = subprocess.Popen(
                [
                    sys.executable,
                    str(server_script),
                    str(project_root / "WORKFLOW.md"),
                    str(port),
                ],
                cwd=project_root,
                env=environment,
                stdout=server_log,
                stderr=subprocess.STDOUT,
                text=True,
            )
            wait_for_api(api, server, root / "server.log")
            handoff = browser_proof.recover_handoff(
                f"http://127.0.0.1:{port}/?pmTab=plan&pmChange={change}#/artifacts",
                handoff_operation,
            )
            assert handoff["state"] == "PUBLISHED"
            managed = api.call("GET", f"/changes/{change}/managed-tickets?project_id={project}")
            assert len(managed["tickets"]) == 4
            managed_item = next(
                value for value in managed["tickets"]
                if value["task_id"] == task_id
            )
            item = {
                "deterministic_ticket_id": managed_item["ticket_identifier"],
                "task_revision": managed_item["task_revision"],
            }
            registry_probe = RunRegistry(project_root / ".symphony" / "state.db")
            mapping = registry_probe.get_pm_managed_ticket(
                project, item["deterministic_ticket_id"]
            )
            recovered_batch = registry_probe.get_pm_handoff_by_operation(
                handoff_operation
            )
            registry_probe.close()
            assert mapping is not None and recovered_batch is not None
            assert recovered_batch["state"] == "PUBLISHED"
            assert recovered_batch["handoff_id"] == mapping.handoff_id
            assert handoff["handoff_id"] == mapping.handoff_id
            handoff_recovery = {
                "staging_response_lost": True,
                "restart_reconciled": True,
                "operation_id": handoff_operation,
                "handoff_id": mapping.handoff_id,
                "staged_item_count": len(staged_items),
                "partial_ticket_count": 0,
                "partial_run_count": 0,
            }
        else:
            approval = api.call("POST", f"/changes/{change}/approvals", {**envelope(project, change, current["revision"]), "confirmation_token": current["approval_confirmation_token"], "reviewer_label": "installed-e2e"})
            current = api.call("GET", f"/changes/{change}?project_id={project}")
            handoff_preview = api.call("POST", f"/changes/{change}/handoffs/preview", envelope(project, change, current["revision"]))
            handoff = api.call("POST", f"/changes/{change}/handoffs", {**envelope(project, change, current["revision"], operation=uuid.uuid4().hex), "preview_token": handoff_preview["preview_token"], "confirmation_token": handoff_preview["preview_token"]})
            assert handoff["state"] == "PUBLISHED"
            item = next(
                value for value in handoff_preview["items"]
                if value["task_id"] == task_id
            )

        ticket = item["deterministic_ticket_id"]
        managed_after_handoff = api.call(
            "GET", f"/changes/{change}/managed-tickets?project_id={project}"
        )["tickets"]
        tickets_by_task = {
            value["task_id"]: value for value in managed_after_handoff
        }
        evidence_tickets = {
            key: tickets_by_task[value]
            for key, value in task_ids.items() if key != "worker"
        }
        reach_ticket = evidence_tickets["native_no_run"]["ticket_identifier"]
        held_tickets = (
            ticket,
            *(
                value["ticket_identifier"]
                for key, value in evidence_tickets.items()
                if key != "native_no_run"
            ),
        )
        set_fixture_paused(project_root, held_tickets, paused=True)
        hold_registry = RunRegistry(project_root / ".symphony" / "state.db")
        try:
            assert all(
                (flags := hold_registry.get_issue_flags(identifier)) is not None
                and flags.paused
                for identifier in held_tickets
            )
            reach_flags = hold_registry.get_issue_flags(reach_ticket)
            assert reach_flags is None or not reach_flags.paused
        finally:
            hold_registry.close()
        reach_ticket_path = project_root / "kanban" / f"{reach_ticket}.md"
        remove_planning_marker(reach_ticket_path)
        workflow_path = project_root / "WORKFLOW.md"
        manual_delivery_workflow = workflow_path.read_text(encoding="utf-8")
        stop_process(server)
        fallback_recovery = exercise_fallback_authority(
            project_root, project, reach_ticket
        )
        conflicting_workflow = manual_delivery_workflow.replace(
            "  auto_merge_on_done: false\n",
            "  auto_merge_on_done: true\n",
            1,
        )
        assert conflicting_workflow != manual_delivery_workflow
        workflow_path.write_text(conflicting_workflow, encoding="utf-8")

        (project_root / ".enable-dispatch").write_text("enabled\n", encoding="utf-8")
        server = subprocess.Popen(
            [sys.executable, str(server_script), str(workflow_path), str(port)],
            cwd=project_root,
            env=environment,
            stdout=server_log,
            stderr=subprocess.STDOUT,
            text=True,
        )
        wait_for_api(api, server, root / "server.log")
        deadline = time.monotonic() + 30
        policy_denial = None
        while time.monotonic() < deadline:
            reach_registry = RunRegistry(project_root / ".symphony" / "state.db")
            try:
                policy_denial = next(
                    (
                        record
                        for record in reach_registry.recent_runs(reach_ticket, limit=10)
                        if record.failure_class == "pm_delivery_policy_conflict"
                    ),
                    None,
                )
            finally:
                reach_registry.close()
            if policy_denial is not None:
                break
            time.sleep(0.1)
        assert policy_denial is not None
        assert policy_denial.backend_agent_pid is None
        assert policy_denial.commit_sha is None
        assert "automatic merge or push" in (policy_denial.failure_message or "")
        reach_registry = RunRegistry(project_root / ".symphony" / "state.db")
        try:
            mapping = reach_registry.get_pm_managed_ticket(project, reach_ticket)
            assert mapping is not None and mapping.handoff_id == handoff["handoff_id"]
            mapping_count = int(reach_registry._connect().execute(
                "SELECT COUNT(*) FROM pm_managed_tickets WHERE project_id = ? AND ticket_identifier = ?",
                (project, reach_ticket),
            ).fetchone()[0])
            assert mapping_count == 1
            denial_records = reach_registry.recent_runs(reach_ticket, limit=10)
            assert len(denial_records) == 1
            assert denial_records[0].failure_class == "pm_delivery_policy_conflict"
            assert denial_records[0].backend_agent_pid is None
            revoked_count = PmAuthority(
                reach_registry, project_root / ".symphony" / "pm-locks"
            ).revoke(project, reach_ticket)
            assert revoked_count == 1
            revocation_count = int(reach_registry._connect().execute(
                """SELECT COUNT(*) FROM pm_authority_revocations
                   WHERE project_id = ? AND change_id = ?
                     AND handoff_id = ? AND task_revision = ?""",
                (
                    project,
                    mapping.change_id,
                    mapping.handoff_id,
                    mapping.task_revision,
                ),
            ).fetchone()[0])
            assert revocation_count == 1
            assert not reach_registry.has_active_pm_dispatch(project, reach_ticket)
        finally:
            reach_registry.close()
        time.sleep(0.75)
        reach_registry = RunRegistry(project_root / ".symphony" / "state.db")
        try:
            denial_records = reach_registry.recent_runs(reach_ticket, limit=10)
            assert len(denial_records) == 1
            assert denial_records[0].failure_class == "pm_delivery_policy_conflict"
            assert denial_records[0].backend_agent_pid is None
            assert not reach_registry.has_active_pm_dispatch(project, reach_ticket)
            reach_flags = reach_registry.get_issue_flags(reach_ticket)
            assert reach_flags is not None and reach_flags.paused
            assert "pm_delivery_policy_conflict" in (
                reach_flags.pause_reason or ""
            )
        finally:
            reach_registry.close()
        combined_reach = {
            **handoff_recovery,
            "server_restarted": True,
            "marker_removed": True,
            "fallback_recognition": fallback_recovery,
            "durable_mapping_count": mapping_count,
            "automatic_delivery_conflict": True,
            "durable_failure_class": policy_denial.failure_class,
            "revocation_record_count": revocation_count,
            "scheduler_denial_record_count": 1,
            "native_process_count": 0,
            "surviving_dispatch_permit": False,
            "dispatch_entry_scope": {
                "initial_attempt_denied_before_backend": True,
                "later_turn_spawn_reached": False,
                "phase_rebuild_reached": False,
                "backend_quota_fallback_reached": False,
                "reason": (
                    "automatic-delivery policy is checked before backend "
                    "construction, so later backend entry points are preempted"
                ),
            },
            "ticket_identifier": reach_ticket,
        }

        stop_process(server)
        workflow_path.write_text(manual_delivery_workflow, encoding="utf-8")
        set_fixture_paused(project_root, (ticket,), paused=False)
        server = subprocess.Popen(
            [sys.executable, str(server_script), str(workflow_path), str(port)],
            cwd=project_root,
            env=environment,
            stdout=server_log,
            stderr=subprocess.STDOUT,
            text=True,
        )
        wait_for_api(api, server, root / "server.log")
        history_url = (
            f"http://127.0.0.1:{port}/?pmTab=idea&pmChange={change}#/artifacts"
        )
        artifact_ref = f"artifact://{ticket}/proof.txt"
        run_id, tested_commit, implementation = await_worker_run(
            project_root, ticket, timeout_seconds=600 if phase_reach else 300
        )
        worker_release = await_worker_release(
            project_root, project, ticket, run_id, tested_commit
        )
        guarded_spawns = [
            json.loads(line)
            for line in (project_root / ".pm-guarded-spawns.jsonl").read_text(encoding="utf-8").splitlines()
            if line.strip()
        ] if (project_root / ".pm-guarded-spawns.jsonl").is_file() else []
        ticket_spawns = [entry for entry in guarded_spawns if entry["ticket"] == ticket and entry["run_id"] == run_id]
        phase_reach_result = None
        if phase_reach:
            # Genuine provider phase rebuild: the Review lane started a new
            # backend process inside the same run, and its launch passed the
            # PM guard again from the live RUNNING permit.
            assert len(ticket_spawns) >= 2, ticket_spawns
            assert ticket_spawns[0]["state_before"] == "RESERVED", ticket_spawns
            assert all(entry["state_before"] == "RUNNING" for entry in ticket_spawns[1:]), ticket_spawns
            assert all(entry["state_after"] == "RUNNING" for entry in ticket_spawns), ticket_spawns
            assert len({entry["pid"] for entry in ticket_spawns}) == len(ticket_spawns), ticket_spawns
            phase_reach_result = {
                "guarded_native_spawns": ticket_spawns,
                "phase_rebuild_reached": True,
                "elapsed_seconds_between_spawns": ticket_spawns[-1]["at"] - ticket_spawns[0]["at"],
                "provider": "codex app-server gpt-6.1-sol (genuine)",
            }

        verification = api.call("POST", f"/changes/{change}/verification", envelope(project, change, payload={"ticket_identifier": ticket, "task_revision": item["task_revision"], "procedure_id": "focused", "run_id": run_id, "artifact_refs": [artifact_ref]}, operation=uuid.uuid4().hex))
        assert verification["result"] == "passed"
        query = urllib.parse.urlencode({"project_id": project, "task_revision": item["task_revision"], "tested_commit": tested_commit})
        acceptance_preview = api.call("GET", f"/tickets/{ticket}/acceptance/preview?{query}")
        assert acceptance_preview["ready"] is True
        if browser_proof is not None:
            browser_proof.accept_verified()
            registry_probe = RunRegistry(project_root / ".symphony" / "state.db")
            acceptance_row = registry_probe.get_active_pm_acceptance(
                project, ticket, item["task_revision"]
            )
            registry_probe.close()
            assert acceptance_row is not None
            acceptance = {"acceptance_id": acceptance_row["acceptance_id"], "state": "ACCEPTED"}
        else:
            acceptance = api.call("POST", f"/tickets/{ticket}/acceptance", {"project_id": project, "task_revision": item["task_revision"], "tested_commit": tested_commit, "confirmation_token": acceptance_preview["confirmation_token"], "reviewer_label": "installed-e2e"})
            assert acceptance["state"] == "ACCEPTED"

        registry_probe = RunRegistry(project_root / ".symphony" / "state.db")
        try:
            assert all(
                not any(
                    record.backend_agent_pid is not None
                    for record in registry_probe.recent_runs(
                        value["ticket_identifier"], limit=10
                    )
                )
                for value in evidence_tickets.values()
            )
        finally:
            registry_probe.close()

        evidence_receipts: dict[str, list[dict]] = {}
        evidence_identifiers = tuple(
            value["ticket_identifier"] for value in evidence_tickets.values()
        )
        evidence_states_before = read_ticket_states(
            project_root, evidence_identifiers
        )
        assert set(evidence_states_before.values()) == {"Todo"}
        reported_start = datetime.now(timezone.utc)
        reported_complete = reported_start + timedelta(seconds=1)

        def submit_no_run(key: str, procedure_id: str) -> dict:
            evidence_item = evidence_tickets[key]
            is_manual = procedure_id == "manual-ux"
            payload = {
                "project_id": project,
                "expected_revision": current["revision"],
                "ticket_identifier": evidence_item["ticket_identifier"],
                "task_revision": evidence_item["task_revision"],
                "tested_commit": tested_commit,
                "procedure_id": procedure_id,
                "environment_summary": (
                    "clean installed disposable fixture; timestamps are "
                    "fixture-reported host observations"
                ),
                "started_at": reported_start.isoformat(),
                "completed_at": reported_complete.isoformat(),
                "result": "passed" if is_manual else "caller-value-ignored",
                "report_ref": (
                    {"ref_kind": "git_blob", "ref": "reports/manual-ux.json"}
                    if is_manual else None
                ),
                "artifact_refs": [],
            }
            preview = api.call(
                "POST", f"/changes/{change}/manual-evidence/preview", payload
            )
            assert preview["ready"] is True, preview
            receipt = api.call(
                "POST", f"/changes/{change}/manual-evidence",
                {
                    **preview["normalized"],
                    "project_id": project,
                    "expected_revision": current["revision"],
                    "client_operation_id": uuid.uuid4().hex,
                    "confirmation_token": preview["confirmation_token"],
                },
            )
            assert receipt["run_id"] is None
            if is_manual:
                assert receipt["procedure_kind"] == "manual"
                assert receipt["return_code"] is None
                download = receipt["report_ref"]["download_url"]
                separator = "&" if "?" in download else "?"
                assert api.read(
                    f"{download.removeprefix('/api/v1/pm')}{separator}project_id={project}"
                ) == b'{"result":"passed","review":"installed"}\n'
            else:
                assert receipt["procedure_kind"] == "native"
                assert receipt["return_code"] == 0
                assert receipt["result"] == "passed"
            evidence_receipts.setdefault(key, []).append(receipt)
            return receipt

        submit_no_run("native_no_run", "focused")
        submit_no_run("manual_only", "manual-ux")
        submit_no_run("mixed", "focused")
        submit_no_run("mixed", "manual-ux")
        evidence_states_after_receipts = read_ticket_states(
            project_root, evidence_identifiers
        )
        assert evidence_states_after_receipts == evidence_states_before

        evidence_acceptances: dict[str, str] = {}
        valid_cross_project_token_rejected = False
        for key, evidence_item in evidence_tickets.items():
            evidence_query = urllib.parse.urlencode({
                "project_id": project,
                "task_revision": evidence_item["task_revision"],
                "tested_commit": tested_commit,
            })
            preview = api.call(
                "GET",
                f"/tickets/{evidence_item['ticket_identifier']}/acceptance/preview?{evidence_query}",
            )
            assert preview["ready"] is True, (key, preview)
            if not valid_cross_project_token_rejected:
                try:
                    api.call(
                        "POST",
                        f"/tickets/{evidence_item['ticket_identifier']}/acceptance",
                        {
                            "project_id": peer_project,
                            "task_revision": evidence_item["task_revision"],
                            "tested_commit": tested_commit,
                            "confirmation_token": preview["confirmation_token"],
                            "reviewer_label": "cross-project-attack",
                        },
                    )
                except RuntimeError as exc:
                    valid_cross_project_token_rejected = (
                        "HTTP 409" in str(exc) and "wrong_project" in str(exc)
                    )
                assert valid_cross_project_token_rejected
            accepted = api.call(
                "POST", f"/tickets/{evidence_item['ticket_identifier']}/acceptance",
                {
                    "project_id": project,
                    "task_revision": evidence_item["task_revision"],
                    "tested_commit": tested_commit,
                    "confirmation_token": preview["confirmation_token"],
                    "reviewer_label": "installed-e2e",
                },
            )
            evidence_acceptances[key] = accepted["acceptance_id"]

        evidence_states_after_acceptance = read_ticket_states(
            project_root, evidence_identifiers
        )
        assert set(evidence_states_after_acceptance.values()) == {"Done"}

        registry_probe = RunRegistry(project_root / ".symphony" / "state.db")
        try:
            assert all(
                not any(
                    record.backend_agent_pid is not None
                    for record in registry_probe.recent_runs(
                        value["ticket_identifier"], limit=10
                    )
                )
                for value in evidence_tickets.values()
            )
            assert all(
                registry_probe.get_accepted_pm_commit(
                    project, value["ticket_identifier"]
                ) == tested_commit
                for value in evidence_tickets.values()
            )
            observed_evidence: dict[str, dict[str, object]] = {}
            for key, value in evidence_tickets.items():
                receipts = registry_probe.list_pm_verification_receipts(
                    project,
                    value["ticket_identifier"],
                    value["task_revision"],
                    tested_commit,
                )
                expected_ids = {
                    receipt["receipt_id"] for receipt in evidence_receipts[key]
                }
                assert {row["receipt_id"] for row in receipts} == expected_ids
                policy = registry_probe.get_pm_required_procedures(
                    project, value["ticket_identifier"]
                )
                assert set(policy) == {
                    receipt["procedure_id"] for receipt in evidence_receipts[key]
                }
                runs = registry_probe.recent_runs(
                    value["ticket_identifier"], limit=10
                )
                observed_evidence[key] = {
                    "scheduler_run_count": len(runs),
                    "native_process_count": sum(
                        record.backend_agent_pid is not None for record in runs
                    ),
                    "receipt_count": len(receipts),
                    "receipt_ids": sorted(expected_ids),
                    "selected_procedure_ids": sorted(policy),
                }
            reach_mapping = registry_probe.get_pm_managed_ticket(
                project, reach_ticket
            )
            assert reach_mapping is not None and reach_mapping.state == "REVOKED"
            persisted_revocations = int(registry_probe._connect().execute(
                """SELECT COUNT(*) FROM pm_authority_revocations
                   WHERE project_id = ? AND change_id = ?
                     AND handoff_id = ? AND task_revision = ?""",
                (
                    project,
                    reach_mapping.change_id,
                    reach_mapping.handoff_id,
                    reach_mapping.task_revision,
                ),
            ).fetchone()[0])
            assert persisted_revocations == 1
            assert not registry_probe.has_active_pm_dispatch(project, reach_ticket)
            state = WorkflowState(project_root / "WORKFLOW.md")
            config, error = state.reload()
            assert config is not None and error is None, error
            reach_issue = FileBoardTracker(config.tracker).fetch_issue_full_by_id(
                reach_ticket
            )
            assert reach_issue is not None and reach_issue.state == "Done"
            reach_authority = PmAuthority(
                registry_probe, project_root / ".symphony" / "pm-locks"
            )
            denied_membership = reach_authority.recognize(reach_issue, project)
            assert denied_membership.managed and not denied_membership.runnable
            try:
                reach_authority.reserve_dispatch(reach_issue, project)
            except PmError as exc:
                assert exc.code == "pm_dispatch_denied"
            else:
                raise AssertionError("revoked accepted ticket regained dispatch authority")
            post_acceptance_native_denial = {
                "board_state": reach_issue.state,
                "mapping_state": reach_mapping.state,
                "revocation_record_count": persisted_revocations,
                "active_permit": False,
                "authority_runnable": False,
                "reserve_denied": True,
            }
        finally:
            registry_probe.close()

        run("git", "cherry-pick", tested_commit, cwd=project_root)
        delivered_commit = run("git", "rev-parse", "HEAD", cwd=project_root)
        tested_tree = run("git", "rev-parse", f"{tested_commit}^{{tree}}", cwd=project_root)
        delivered_tree = run("git", "rev-parse", f"{delivered_commit}^{{tree}}", cwd=project_root)
        assert delivered_tree == tested_tree, "local delivery tree differs from the tested worker tree"
        assert (project_root / "implemented.txt").read_text(encoding="utf-8") == "verified\n"

        if browser_proof is not None:
            browser_proof.archive()
            archived = sorted((planning / "openspec" / "changes" / "archive").glob("*-installed-proof"))
            assert len(archived) == 1
            archive = {"archived_path": str(archived[0]), "baseline_paths": [], "output": "browser-confirmed"}
        else:
            archive_preview = api.call("POST", f"/changes/{change}/archive/preview", envelope(project, change, current["revision"]))
            archive = api.call("POST", f"/changes/{change}/archive/{archive_preview['archive_id']}", {**envelope(project, change, current["revision"], operation=uuid.uuid4().hex), "confirmation_token": archive_preview["confirmation_token"]})
        archived_path = Path(archive["archived_path"])
        assert archived_path.is_dir() and planning in archived_path.parents
        archived_change = api.call(
            "GET", f"/changes/{change}?project_id={project}"
        )
        assert archived_change["change_id"] == change
        assert archived_change["state"] == "Archived"
        assert archived_change["workspace_state"] == "RELEASED"
        assert archived_change["approval_confirmation_token"] is None
        archived_document = api.call(
            "GET", f"/changes/{change}/documents/tasks?project_id={project}"
        )
        assert archived_document["editable"] is False
        assert archived_document["unsupported_reason"] == "archived planning history"
        archived_wbs = api.call(
            "GET", f"/changes/{change}/wbs?project_id={project}"
        )
        assert task_id in {value["task_id"] for value in archived_wbs["tasks"]}
        archived_tickets = api.call(
            "GET", f"/changes/{change}/managed-tickets?project_id={project}"
        )
        assert ticket in {
            value["ticket_identifier"] for value in archived_tickets["tickets"]
        }
        archive_denials: dict[str, bool] = {}
        denied_calls = {
            "patch": (
                "PATCH", f"/changes/{change}/documents/proposal",
                envelope(
                    project, change, archived_change["revision"],
                    {"operations": [{
                        "op": "replace_source", "source": "# forbidden\n",
                    }]},
                ),
            ),
            "approval": (
                "POST", f"/changes/{change}/approvals",
                {**envelope(project, change, archived_change["revision"]),
                 "confirmation_token": "forbidden", "reviewer_label": "forbidden"},
            ),
            "handoff": (
                "POST", f"/changes/{change}/handoffs/preview",
                envelope(project, change, archived_change["revision"]),
            ),
            "reclaim": (
                "POST", f"/changes/{change}/workspace/claim",
                {"project_id": project, "change_id": change,
                 "workspace_path": str(planning)},
            ),
        }
        for name, (method, path, body) in denied_calls.items():
            try:
                api.call(method, path, body)
            except RuntimeError as exc:
                archive_denials[name] = "HTTP 409" in str(exc)
            else:
                archive_denials[name] = False
        assert all(archive_denials.values()), archive_denials
        if browser_proof is not None:
            browser_proof.review_archived_history(history_url, task_id)
        return {
            "phase_reach": phase_reach_result,
            "root": str(root),
            "url": f"http://127.0.0.1:{port}/#/artifacts",
            "change_id": change,
            "validation": validation,
            "approval_id": approval["approval_id"],
            "handoff_id": handoff["handoff_id"],
            "ticket": ticket,
            "tested_commit": tested_commit,
            "verification_receipt": verification["receipt_id"],
            "acceptance_id": acceptance["acceptance_id"],
            "post_acceptance_native_denial": post_acceptance_native_denial,
            "evidence_modes": {
                key: {
                    "ticket_identifier": evidence_tickets[key]["ticket_identifier"],
                    "receipt_ids": [value["receipt_id"] for value in values],
                    "procedure_kinds": [value["procedure_kind"] for value in values],
                    "return_codes": [value["return_code"] for value in values],
                    "acceptance_id": evidence_acceptances[key],
                    **observed_evidence[key],
                    "state_before_evidence": evidence_states_before[
                        evidence_tickets[key]["ticket_identifier"]
                    ],
                    "state_after_evidence": evidence_states_after_receipts[
                        evidence_tickets[key]["ticket_identifier"]
                    ],
                    "state_after_acceptance": evidence_states_after_acceptance[
                        evidence_tickets[key]["ticket_identifier"]
                    ],
                    "timestamp_source": "fixture-reported host observations",
                }
                for key, values in evidence_receipts.items()
            },
            "archive": archive,
            "archive_history": {
                "original_url": history_url,
                "state": archived_change["state"],
                "document_ids": sorted(
                    value["document_id"] for value in archived_change["documents"]
                ),
                "task_ids": sorted(
                    value["task_id"] for value in archived_wbs["tasks"]
                ),
                "ticket_ids": sorted(
                    value["ticket_identifier"]
                    for value in archived_tickets["tickets"]
                ),
                "artifact_refs": [artifact_ref],
                "denied_mutations": archive_denials,
            },
            "assistance": assistance_receipt,
            "browser": browser_proof.close() if browser_proof is not None else {"requested": False},
            "execution": {
                "scheduler_dispatched": True,
                "agent_kind": "codex",
                "model": "gpt-6.1-sol",
                "effort": "medium",
                "run_id": run_id,
                "workspace": str(implementation),
                "release": worker_release,
                "doctor": doctor.stdout,
            },
            "combined_recovery_reach": combined_reach,
            "delivery": {
                "mode": "local-disposable-cherry-pick",
                "source_commit": tested_commit,
                "delivered_commit": delivered_commit,
                "tree_sha": delivered_tree,
                "remote_count": len(run("git", "remote", cwd=project_root).splitlines()),
            },
            "boundaries": {
                "openspec_validation": "actual-cli",
                "handoff": "native-file-tracker",
                "verification": "host-executed-approved-command",
                "acceptance": "human-confirmation-token",
                "archive": "journaled-production-adapter",
                "raw_work_item_identity_attack_rejected": raw_identity_rejected,
                "client_change_identity_rejected": client_change_id_rejected,
                "same_slug_registered_project_refused": same_slug_stale_refused,
                "valid_cross_project_acceptance_token_rejected": valid_cross_project_token_rejected,
                "registered_project_id": project,
                "peer_registered_project_id": peer_project,
                "accepted_baseline_requirement_id": requirement_id,
                "accepted_baseline_scenario_id": scenario_id,
                "replaced_fresh_requirement_id": fresh_requirement_id,
                "retained_fresh_scenario_id": fresh_scenario_id,
            },
        }
    finally:
        stop_process(server)
        if browser_proof is not None and browser_proof.browser.is_connected():
            browser_proof.close()
        server_log.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--openspec", type=Path, required=True)
    parser.add_argument("--port", type=int, default=18942)
    parser.add_argument("--assistance", action="store_true")
    parser.add_argument("--browser", action="store_true")
    parser.add_argument("--phase-reach", action="store_true")
    parser.add_argument("--receipt", type=Path)
    parser.add_argument("--wheel", type=Path)
    parser.add_argument("--source-sha")
    parser.add_argument("--expected-prefix", type=Path)
    args = parser.parse_args()
    if args.root.exists() and any(args.root.iterdir()):
        raise SystemExit("--root must be absent or empty")
    args.root.mkdir(parents=True, exist_ok=True)
    module_path = Path(symphony.pm.__file__).resolve()
    if args.expected_prefix is not None:
        prefix = args.expected_prefix.resolve()
        if prefix != module_path and prefix not in module_path.parents:
            raise SystemExit(f"installed package escaped expected prefix: {module_path}")
    receipt = journey(
        args.root.resolve(), args.openspec.resolve(), args.port,
        assistance=args.assistance, browser=args.browser, phase_reach=args.phase_reach,
    )
    receipt["installation"] = {
        "python": sys.executable,
        "module_path": str(module_path),
        "source_sha": args.source_sha,
        "wheel": str(args.wheel.resolve()) if args.wheel else None,
        "wheel_sha256": sha256(args.wheel.resolve()) if args.wheel else None,
        "distributions": distribution_inventory(),
        "openspec_executable": str(args.openspec.resolve()),
        "openspec_version": run(str(args.openspec.resolve()), "--version", cwd=args.root.resolve()),
    }
    destination = args.receipt or args.root / "receipt.json"
    destination.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(destination)


if __name__ == "__main__":
    main()
