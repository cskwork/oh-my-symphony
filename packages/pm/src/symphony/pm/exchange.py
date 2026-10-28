"""Constrained XLSX/CSV exchange with trusted-base three-way comparison."""

from __future__ import annotations

import csv
import hashlib
import io
import posixpath
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from typing import Any, Literal, cast
from xml.etree import ElementTree

from .models import (
    ChangeView,
    ExportPayload,
    ImportCandidate,
    ImportFieldChange,
    PmError,
    PmFinding,
    WbsDependency,
    WbsTask,
)
from .wbs import parse_wbs, validate_wbs_limits

MAX_COMPRESSED = 5 * 1024 * 1024
MAX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ENTRIES = 1_000
MAX_TASKS = 5_000
MAX_EDGES = 20_000
MAX_TEXT = 16 * 1024
REQUIRED_SHEETS = ("WBS", "Metadata")
OPTIONAL_EDIT_SHEETS = ("Task Details", "Dependencies", "Traceability")
_FORMULA_PREFIX = ("=", "+", "-", "@")
_FORBIDDEN_PARTS = (
    "vbaproject.bin",
    "activex",
    "embeddings/",
    "oleobject",
)


@dataclass(frozen=True)
class WorkbookExchange:
    max_workbook_bytes: int = MAX_COMPRESSED
    max_uncompressed_bytes: int = MAX_UNCOMPRESSED
    max_zip_entries: int = MAX_ENTRIES
    max_nodes: int = MAX_TASKS
    max_edges: int = MAX_EDGES
    max_text_chars: int = MAX_TEXT

    def export(
        self,
        change_snapshot: ChangeView,
        runtime_projection: dict[str, Any],
        *,
        export_id: str | None = None,
    ) -> ExportPayload:
        openpyxl = _openpyxl()
        tasks = _tasks(change_snapshot)
        self._validate_tasks(tasks)
        workbook = openpyxl.Workbook()
        active = workbook.active
        if active is not None:
            workbook.remove(active)
        display_numbers = _display_numbers(tasks, self.max_text_chars)
        titles = {task.task_id: task.title for task in tasks}
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        readme = workbook.create_sheet("Read Me")
        wbs = workbook.create_sheet("WBS")
        details = workbook.create_sheet("Task Details")
        deps = workbook.create_sheet("Dependencies")
        trace = workbook.create_sheet("Traceability")
        metadata = workbook.create_sheet("Metadata")

        _append_literal(readme, ("Project", change_snapshot.project_id))
        _append_literal(readme, ("Change", change_snapshot.change_id))
        _append_literal(readme, ("Revision", change_snapshot.revision))
        _append_literal(readme, ("State", change_snapshot.state))
        _append_literal(
            readme,
            (
                "Editing",
                "Existing Task ID and runtime columns are read-only. Use Action=Add/Delete explicitly.",
            ),
        )
        _append_literal(readme, ("Exported at (UTC)", datetime.now(UTC).isoformat()))
        _append_literal(readme, (
            "Execution snapshots",
            "Execution Snapshot values and the leaf status summary are export-time snapshots. "
            "They do not update in this workbook and are ignored during import.",
        ))
        _append_literal(readme, (
            "Status legend",
            "Supplied statuses use the runtime state's own name. Not supplied means no execution "
            "state was supplied; the planning State above is separate. Up to 48 distinct supplied "
            "statuses are listed; additional statuses are counted together. Full values remain in WBS.",
        ))
        _append_literal(readme, (
            "Read-only shading",
            "Gray cells identify existing Task IDs, derived display/dependency fields, runtime "
            "links/snapshots and export provenance. "
            "For Action=Add, enter a NEW-* temporary Task ID. Other task fields remain editable.",
        ))
        _append_literal(readme, (
            "Dropdown hints",
            "Dropdowns suggest existing Action, dependency Kind and Reference Type values. "
            "Import preview validates edits. Type, Priority and Agent Override accept custom values.",
        ))
        _append_literal(readme, (
            "Defaults",
            "Blank Action means Keep. Blank dependency Kind means requires_code. "
            "Reference Type is case-insensitive.",
        ))
        parents = {task.task_id: task.parent_id for task in tasks}
        parent_ids = {task.parent_id for task in tasks if task.parent_id}
        leaves = [task for task in tasks if task.kind != "Group" and task.task_id not in parent_ids]
        statuses = Counter(
            str(runtime_projection.get(task.task_id, {}).get("state") or "")
            for task in leaves
        )
        _append_literal(readme, ("Leaf tasks", len(leaves)))
        _append_literal(readme, ("Leaf status summary", "Leaf task count"))
        _append_literal(readme, ("Not supplied", statuses.pop("", 0)))
        supplied_statuses = sorted(statuses.items())
        for state, count in supplied_statuses[:48]:
            _append_literal(readme, (f"Supplied: {state}", count))
        if len(supplied_statuses) > 48:
            _append_literal(readme, (
                "Additional supplied statuses", sum(count for _, count in supplied_statuses[48:]),
            ))

        _append_literal(
            wbs,
            (
                "Action",
                "Task ID",
                "Parent ID",
                "Order",
                "Type",
                "Title",
                "Owner",
                "Agent Override",
                "Priority",
                "Target Date",
                "Estimate",
                "Execution Snapshot",
                "Display Number",
                "Dependency Summary",
            ),
        )
        _append_literal(
            details,
            (
                "Task ID",
                "Description",
                "Acceptance",
                "Verification",
                "Integration Impact",
                "Risk",
                "Notes",
            ),
        )
        _append_literal(
            deps, ("Task ID", "Predecessor ID", "Kind", "Justification", "Action", "Task Title", "Predecessor Title")
        )
        _append_literal(
            trace,
            ("Task ID", "Reference Type", "Reference ID", "Runtime Link", "Action"),
        )
        _append_literal(
            metadata,
            (
                "format_version",
                "project_id",
                "change_id",
                "revision",
                "scope_hash",
                "export_id",
            ),
        )
        _append_literal(
            metadata,
            (
                "1",
                change_snapshot.project_id,
                change_snapshot.change_id,
                change_snapshot.revision,
                change_snapshot.scope_hash,
                export_id or "",
            ),
        )
        for task in tasks:
            runtime = runtime_projection.get(task.task_id, {})
            _append_literal(
                wbs,
                (
                    "Keep",
                    task.task_id,
                    task.parent_id or "",
                    task.order,
                    task.kind,
                    task.title,
                    task.human_owner or "",
                    task.agent_kind_override or "",
                    task.priority or "",
                    task.target_date or "",
                    task.estimate or "",
                    str(runtime.get("state") or ""),
                    display_numbers.get(task.task_id, ""),
                    _bounded_derived("\n".join(
                        f"{dependency.task_id}: {titles.get(dependency.task_id, '')} ({dependency.kind})"
                        for dependency in task.dependencies
                    ), self.max_text_chars, "Dependencies"),
                ),
            )
            _append_literal(
                details,
                (
                    task.task_id,
                    task.description,
                    "\n".join(task.acceptance),
                    "\n".join(task.verification),
                    task.integration_impact,
                    task.risk,
                    task.notes,
                ),
            )
            for dependency in task.dependencies:
                _append_literal(
                    deps,
                    (
                        task.task_id,
                        dependency.task_id,
                        dependency.kind,
                        dependency.justification or "",
                        "Keep",
                        task.title,
                        titles.get(dependency.task_id, ""),
                    ),
                )
            for reference in task.requirement_refs:
                _append_literal(trace, (task.task_id, "scenario" if reference.startswith("SC-") else "requirement", reference, str(runtime.get("link") or ""), "Keep"))
            for kind, references in (("storyboard", task.storyboard_refs), ("design", task.design_refs)):
                for reference in references:
                    _append_literal(trace, (task.task_id, kind, reference, str(runtime.get("link") or ""), "Keep"))
        frozen = {"WBS": "D2", "Task Details": "B2", "Dependencies": "C2", "Traceability": "B2"}
        read_only_columns = {
            "Read Me": {1, 2}, "WBS": {2, 12, 13, 14}, "Task Details": {1},
            "Dependencies": {1, 6, 7}, "Traceability": {1, 4}, "Metadata": set(range(1, 7)),
        }
        read_only_fill = PatternFill("solid", fgColor="E9EDF2")
        for sheet in workbook.worksheets:
            sheet.freeze_panes = frozen.get(sheet.title, "A2")
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                values = [str(cell.value or "") for cell in column[:100]]
                width = min(60, max(12, max((len(value) for value in values), default=0) + 2))
                column_index = column[0].column
                if not isinstance(column_index, int):
                    continue
                sheet.column_dimensions[get_column_letter(column_index)].width = width
                for cell in column:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if column_index in read_only_columns[sheet.title]:
                        cell.fill = read_only_fill
                        cell.font = Font(color="475569")
                    if cell.row == 1 and sheet is not readme:
                        cell.fill = PatternFill("solid", fgColor=(
                            "475569" if column_index in read_only_columns[sheet.title] else "1E3A5F"
                        ))
                        cell.font = Font(color="FFFFFF", bold=True)
        for row, task in enumerate(tasks, 2):
            depth = 0
            parent = task.parent_id
            seen = {task.task_id}
            while parent in parents and parent not in seen and depth < 250:
                seen.add(parent)
                depth += 1
                parent = parents[parent]
            wbs.cell(row, 6).alignment = Alignment(vertical="top", wrap_text=True, indent=depth)
            wbs.row_dimensions[row].outlineLevel = min(depth, 7)
            wbs.cell(row, 10).number_format = "yyyy-mm-dd"
        for sheet, column, values, limit, prompt in (
            (wbs, "A", "Keep,Add,Delete", self.max_nodes, "Blank means Keep. Add uses a NEW-* Task ID."),
            (deps, "C", "requires_code,order_only", self.max_edges, "Blank means requires_code."),
            (deps, "E", "Keep,Add,Delete", self.max_edges, "Blank means Keep."),
            (trace, "B", "requirement,scenario,storyboard,design,api", self.max_edges, "Reference Type is case-insensitive."),
            (trace, "E", "Keep,Add,Delete", self.max_edges, "Blank means Keep."),
        ):
            validation = DataValidation(
                type="list", formula1=f'"{values}"', allow_blank=True,
                showDropDown=False, showErrorMessage=False, showInputMessage=True,
                promptTitle="Import value hint", prompt=prompt,
            )
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}{max(2, limit + 1)}")
        output = io.BytesIO()
        workbook.save(output)
        data = output.getvalue()
        _preflight(
            data, max_compressed=self.max_workbook_bytes,
            max_uncompressed=self.max_uncompressed_bytes,
            max_entries=self.max_zip_entries,
        )
        filename = f"{change_snapshot.change_id}_wbs_{change_snapshot.revision[7:19]}.xlsx"
        return ExportPayload(
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data,
            {
                "project_id": change_snapshot.project_id,
                "change_id": change_snapshot.change_id,
                "revision": change_snapshot.revision,
                "scope_hash": change_snapshot.scope_hash,
                "file_hash": _hash(data),
                "export_id": export_id or "",
            },
        )

    def inspect_metadata(self, file_bytes: bytes) -> dict[str, str]:
        """Read the small literal metadata row after the same hostile ZIP checks."""
        _preflight(
            file_bytes, max_compressed=self.max_workbook_bytes,
            max_uncompressed=self.max_uncompressed_bytes,
            max_entries=self.max_zip_entries,
        )
        openpyxl = _openpyxl()
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(file_bytes), read_only=True, data_only=False,
                keep_links=False, keep_vba=False,
            )
        except Exception as exc:
            raise PmError("unsafe_workbook", f"workbook cannot be opened: {exc}") from exc
        try:
            if "Metadata" not in workbook.sheetnames:
                raise PmError("unsafe_workbook", "required workbook sheets are missing")
            sheet = workbook["Metadata"]
            if sheet.max_row > 2 or sheet.max_column > 6:
                raise PmError("unsafe_workbook", "worksheet dimensions exceed limit: Metadata")
            records = _sheet_records(sheet, max_text_chars=self.max_text_chars)
            if len(records) != 1:
                raise PmError("unsafe_workbook", "metadata sheet is invalid")
            return {str(key): str(value or "") for key, value in records[0].items()}
        finally:
            workbook.close()

    def preview_import(
        self,
        file_bytes: bytes,
        trusted_export_snapshot: ChangeView,
        current_change: ChangeView,
        *,
        filename: str = "workbook.xlsx",
    ) -> ImportCandidate:
        _preflight(
            file_bytes, max_compressed=self.max_workbook_bytes,
            max_uncompressed=self.max_uncompressed_bytes,
            max_entries=self.max_zip_entries,
        )
        openpyxl = _openpyxl()
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(file_bytes),
                read_only=True,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
        except Exception as exc:
            raise PmError("unsafe_workbook", f"workbook cannot be opened: {exc}") from exc
        try:
            missing = sorted(set(REQUIRED_SHEETS) - set(workbook.sheetnames))
            if missing:
                raise PmError(
                    "unsafe_workbook",
                    "required workbook sheets are missing",
                    details={"missing": missing},
                )
            for sheet in workbook.worksheets:
                limit = self._sheet_limits().get(sheet.title)
                if limit is None:
                    raise PmError("unsafe_workbook", f"unexpected sheet {sheet.title}")
                if sheet.max_row > limit[0] or sheet.max_column > limit[1]:
                    raise PmError(
                        "unsafe_workbook",
                        f"worksheet dimensions exceed limit: {sheet.title}",
                    )
                if hasattr(sheet, "reset_dimensions"):
                    sheet.reset_dimensions()
                    cast(Any, sheet).calculate_dimension(force=True)
                    if sheet.max_row > limit[0] or sheet.max_column > limit[1]:
                        raise PmError(
                            "unsafe_workbook",
                            f"worksheet dimensions exceed limit: {sheet.title}",
                        )
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            raise PmError(
                                "unsafe_workbook",
                                f"formula cell: {sheet.title}!{cell.coordinate}",
                            )
            metadata = _sheet_records(
                workbook["Metadata"], max_text_chars=self.max_text_chars
            )
            if len(metadata) != 1:
                raise PmError("unsafe_workbook", "metadata sheet is invalid")
            meta = metadata[0]
            if str(meta.get("format_version") or "") != "1":
                raise PmError("unsafe_workbook", "unsupported workbook format_version")
            expected = {
                "project_id": trusted_export_snapshot.project_id,
                "change_id": trusted_export_snapshot.change_id,
                "revision": trusted_export_snapshot.revision,
                "scope_hash": trusted_export_snapshot.scope_hash,
            }
            for key, value in expected.items():
                if str(meta.get(key) or "") != value:
                    raise PmError(
                        "revision_conflict",
                        f"workbook {key} does not match trusted export",
                        status=409,
                    )
            base = {task.task_id: task for task in _tasks(trusted_export_snapshot)}
            current = {task.task_id: task for task in _tasks(current_change)}
            self._validate_tasks(tuple(base.values()))
            self._validate_tasks(tuple(current.values()))
            rows = _sheet_records(
                workbook["WBS"], max_text_chars=self.max_text_chars
            )
            if len(rows) > self.max_nodes:
                raise PmError("unsafe_workbook", "workbook task limit exceeded")
            changes: list[ImportFieldChange] = []
            conflicts: list[ImportFieldChange] = []
            findings: list[PmFinding] = []
            assigned: dict[str, str] = {}
            seen: set[str] = set()
            imported_rows: dict[str, dict[str, Any]] = {}
            for row in rows:
                task_id = str(row.get("Task ID") or "").strip()
                action = str(row.get("Action") or "Keep").strip()
                if action not in {"Keep", "Add", "Delete"}:
                    raise PmError("unsafe_workbook", f"unknown WBS action {action}")
                if not task_id:
                    raise PmError("unsafe_workbook", "task row has no ID")
                if task_id in seen:
                    raise PmError("unsafe_workbook", f"duplicate task ID {task_id}")
                seen.add(task_id)
                if action == "Add":
                    if not task_id.startswith("NEW-"):
                        raise PmError(
                            "unsafe_workbook", "new rows require a NEW-* temporary ID"
                        )
                    assigned[task_id] = _permanent_id(
                        current_change.change_id,
                        task_id,
                        str(row.get("Type") or "Feature") == "Group",
                    )
                elif task_id not in base:
                    raise PmError(
                        "unsafe_workbook", f"unknown existing task ID {task_id}"
                    )
                imported_rows[task_id] = row
            details_present = "Task Details" in workbook.sheetnames
            dependencies_present = "Dependencies" in workbook.sheetnames
            traceability_present = "Traceability" in workbook.sheetnames
            details = _records_by_id(
                workbook["Task Details"], "Task ID", self.max_text_chars
            ) if details_present else {}
            dependencies = _dependency_rows(
                workbook["Dependencies"], self.max_text_chars, self.max_edges
            ) if dependencies_present else {}
            traceability = _trace_rows(
                workbook["Traceability"], self.max_text_chars
            ) if traceability_present else {}
            known_ids = set(base) | set(assigned)
            for sheet_name, grouped in (
                ("Task Details", details),
                ("Dependencies", dependencies),
                ("Traceability", traceability),
            ):
                unknown = sorted(set(grouped) - known_ids)
                if unknown:
                    raise PmError(
                        "unsafe_workbook",
                        f"{sheet_name} references unknown task ID {unknown[0]}",
                    )
            for task_id, row in imported_rows.items():
                action = str(row.get("Action") or "Keep").strip()
                resolved_id = assigned.get(task_id, task_id)
                if action == "Delete":
                    if _delete_scope_changed(task_id, base, current):
                        conflicts.append(ImportFieldChange(task_id, "__delete__", _task_snapshot(base[task_id]), _task_snapshot(current.get(task_id)), True, "delete"))
                    else:
                        changes.append(ImportFieldChange(task_id, "__delete__", _task_snapshot(base[task_id]), _task_snapshot(current.get(task_id)), True, "delete"))
                    continue
                baseline = base.get(task_id)
                imported_fields = _wbs_fields(row, baseline, assigned)
                imported_fields.update(_detail_fields(details.get(task_id), baseline))
                if dependencies_present and task_id in dependencies:
                    imported_fields["dependencies"] = _resolve_dependencies(
                        dependencies[task_id], assigned, known_ids, findings, task_id,
                        baseline.dependencies if baseline is not None else (),
                    )
                if traceability_present and task_id in traceability:
                    imported_fields.update(_resolve_traceability(
                        traceability[task_id], findings, task_id, baseline
                    ))
                if action == "Add":
                    imported_fields["task_id"] = resolved_id
                    changes.append(ImportFieldChange(resolved_id, "__add__", None, None, imported_fields, "add"))
                    continue
                assert baseline is not None
                current_task = current.get(task_id)
                if current_task is None:
                    findings.append(PmFinding("missing_current_task", f"task {task_id} no longer exists", severity="blocking", item_id=task_id))
                    continue
                for field, imported in imported_fields.items():
                    base_value = getattr(baseline, field)
                    current_value = getattr(current_task, field)
                    if imported == base_value:
                        continue
                    change = ImportFieldChange(task_id, field, base_value, current_value, imported, "set")
                    if current_value != base_value and current_value != imported:
                        conflicts.append(change)
                    else:
                        changes.append(change)
            digest = hashlib.sha256()
            digest.update(_hash(file_bytes).encode())
            digest.update(filename.encode("utf-8"))
            digest.update(current_change.revision.encode())
            digest.update(repr(changes).encode())
            digest.update(repr(conflicts).encode())
            return ImportCandidate(
                _hash(file_bytes),
                filename,
                trusted_export_snapshot.revision,
                current_change.revision,
                "sha256:" + digest.hexdigest(),
                tuple(changes),
                tuple(conflicts),
                assigned,
                tuple(findings),
            )
        finally:
            workbook.close()

    def render_csv(
        self, change_snapshot: ChangeView, *, excel_utf8: bool
    ) -> bytes:
        tasks = _tasks(change_snapshot)
        self._validate_tasks(tasks)
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(("Task ID", "Parent ID", "Order", "Type", "Title"))
        for task in tasks:
            writer.writerow(
                tuple(
                    _csv_safe(str(value or ""))
                    for value in (
                        task.task_id,
                        task.parent_id,
                        task.order,
                        task.kind,
                        task.title,
                    )
                )
            )
        encoding = "utf-8-sig" if excel_utf8 else "utf-8"
        return output.getvalue().encode(encoding)

    def _sheet_limits(self) -> dict[str, tuple[int, int]]:
        return {
            "Read Me": (64, 2),
            "WBS": (self.max_nodes + 1, 14),
            "Task Details": (self.max_nodes + 1, 7),
            "Dependencies": (self.max_edges + 1, 7),
            "Traceability": (self.max_edges + 1, 5),
            "Metadata": (2, 6),
        }

    def _validate_tasks(self, tasks: tuple[WbsTask, ...]) -> None:
        findings = validate_wbs_limits(
            tasks, max_nodes=self.max_nodes, max_edges=self.max_edges,
            max_text_chars=self.max_text_chars,
        )
        if findings:
            raise PmError(
                "unsafe_workbook", "planning data exceeds configured exchange limits",
                details={"findings": [item.code for item in findings]},
            )


def _preflight(
    data: bytes, *, max_compressed: int, max_uncompressed: int,
    max_entries: int,
) -> None:
    if len(data) > max_compressed:
        raise PmError("unsafe_workbook", "compressed workbook limit exceeded")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise PmError("unsafe_workbook", "file is not a valid XLSX container") from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > max_entries:
            raise PmError("unsafe_workbook", "ZIP entry limit exceeded")
        total = 0
        for info in infos:
            normalized = posixpath.normpath(info.filename)
            if (
                normalized.startswith("../")
                or normalized.startswith("/")
                or "\\" in normalized
            ):
                raise PmError("unsafe_workbook", "unsafe ZIP path")
            total += info.file_size
            if total > max_uncompressed:
                raise PmError("unsafe_workbook", "uncompressed-size limit")
            lower = normalized.lower()
            if any(part in lower for part in _FORBIDDEN_PARTS):
                raise PmError("unsafe_workbook", f"forbidden workbook part: {normalized}")
            if lower.endswith((".xml", ".rels")):
                xml = archive.read(info)
                folded = xml.lower()
                if b"<!doctype" in folded or b"<!entity" in folded:
                    raise PmError(
                        "unsafe_workbook",
                        f"DTD or entity declaration in workbook XML: {normalized}",
                    )
                try:
                    root = ElementTree.fromstring(xml)
                except (ElementTree.ParseError, ValueError) as exc:
                    raise PmError(
                        "unsafe_workbook",
                        f"malformed workbook XML: {normalized}",
                    ) from exc
            else:
                root = None
            if lower.endswith(".rels") and root is not None:
                for relationship in root:
                    if relationship.attrib.get("TargetMode") == "External":
                        raise PmError("unsafe_workbook", "external relationship")


def _bounded_derived(value: str, limit: int, sheet: str) -> str:
    if len(value) <= limit:
        return value
    notice = f"... See {sheet}."
    return value[:max(0, limit - len(notice))] + notice[:limit]


def _display_numbers(tasks: tuple[WbsTask, ...], limit: int) -> dict[str, str]:
    task_ids = {task.task_id for task in tasks}
    children: dict[str | None, list[WbsTask]] = {}
    for task in tasks:
        parent = task.parent_id if task.parent_id in task_ids else None
        children.setdefault(parent, []).append(task)
    for siblings in children.values():
        siblings.sort(key=lambda task: (task.order, task.task_id))
    stack = [(task, str(index)) for index, task in reversed(list(enumerate(children.get(None, []), 1)))]
    numbers = {}
    while stack:
        task, number = stack.pop()
        if task.task_id in numbers:
            continue
        numbers[task.task_id] = _bounded_derived(number, limit, "Parent ID")
        stack.extend(
            (child, f"{number}.{index}")
            for index, child in reversed(list(enumerate(children.get(task.task_id, []), 1)))
        )
    return numbers


def _tasks(change: ChangeView) -> tuple[WbsTask, ...]:
    document = next((item for item in change.documents if item.document_id == "tasks"), None)
    if document is None:
        return ()
    parsed = parse_wbs(document.source)
    blocking = [item for item in parsed.findings if item.severity == "blocking"]
    if blocking:
        raise PmError(
            "unsupported_format",
            "tasks.md is not safe for exchange",
            details={"findings": [item.code for item in blocking]},
        )
    return parsed.tasks


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise PmError(
            "capability_unavailable",
            "XLSX exchange requires the pm-exchange optional dependencies",
            status=503,
        ) from exc
    from openpyxl.xml import DEFUSEDXML

    if not DEFUSEDXML:
        raise PmError(
            "capability_unavailable", "openpyxl is not using defusedxml", status=503
        )
    return openpyxl


def _append_literal(sheet: Any, values: tuple[Any, ...]) -> None:
    sheet.append(list(values))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _sheet_records(sheet: Any, *, max_text_chars: int) -> list[dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return []
    headers = [str(item or "") for item in raw_headers]
    records: list[dict[str, Any]] = []
    for values in iterator:
        if not any(value not in (None, "") for value in values):
            continue
        record = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
        for value in record.values():
            if isinstance(value, str) and len(value) > max_text_chars:
                raise PmError("unsafe_workbook", "workbook text field limit exceeded")
        records.append(record)
    return records


def _records_by_id(
    sheet: Any, id_column: str, max_text_chars: int
) -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for record in _sheet_records(sheet, max_text_chars=max_text_chars):
        item_id = str(record.get(id_column) or "").strip()
        if not item_id:
            raise PmError("unsafe_workbook", f"{sheet.title} row has no {id_column}")
        if item_id in records:
            raise PmError("unsafe_workbook", f"duplicate {id_column} {item_id}")
        records[item_id] = record
    return records


def _wbs_fields(
    row: dict[str, Any], baseline: WbsTask | None, assigned: dict[str, str]
) -> dict[str, Any]:
    default_order = baseline.order if baseline is not None else 0
    default_kind = baseline.kind if baseline is not None else "Feature"
    parent = _none(row.get("Parent ID"))
    return {
        "parent_id": assigned.get(parent, parent) if parent else None,
        "order": _integer(row.get("Order"), default_order),
        "kind": str(row.get("Type") or default_kind),
        "title": str(row.get("Title") or ""),
        "human_owner": _none(row.get("Owner")),
        "agent_kind_override": _none(row.get("Agent Override")),
        "priority": _none(row.get("Priority")),
        "target_date": _none(row.get("Target Date")),
        "estimate": _none(row.get("Estimate")),
    }


def _detail_fields(
    row: dict[str, Any] | None, baseline: WbsTask | None
) -> dict[str, Any]:
    if row is None:
        return {}
    return {
        "description": str(row.get("Description") or ""),
        "acceptance": _lines(row.get("Acceptance")),
        "verification": _lines(row.get("Verification")),
        "integration_impact": str(row.get("Integration Impact") or ""),
        "risk": str(row.get("Risk") or ""),
        "notes": str(row.get("Notes") or ""),
    }


def _dependency_rows(
    sheet: Any, max_text_chars: int, max_edges: int
) -> dict[str, tuple[dict[str, Any], ...]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in _sheet_records(sheet, max_text_chars=max_text_chars):
        task_id = str(row.get("Task ID") or "").strip()
        if not task_id:
            raise PmError("unsafe_workbook", "dependency row has no Task ID")
        grouped.setdefault(task_id, []).append(row)
    if sum(len(rows) for rows in grouped.values()) > max_edges:
        raise PmError("unsafe_workbook", "workbook dependency limit exceeded")
    return {key: tuple(value) for key, value in grouped.items()}


def _resolve_dependencies(
    rows: tuple[dict[str, Any], ...],
    assigned: dict[str, str],
    known_ids: set[str],
    findings: list[PmFinding],
    item_id: str,
    baseline: tuple[WbsDependency, ...],
) -> tuple[WbsDependency, ...]:
    values: list[WbsDependency] = list(baseline)
    for row in rows:
        predecessor = str(row.get("Predecessor ID") or "").strip()
        kind = str(row.get("Kind") or "requires_code").strip()
        action = str(row.get("Action") or "Keep").strip()
        if action not in {"Keep", "Add", "Delete"} or predecessor not in known_ids or kind not in {"requires_code", "order_only"}:
            findings.append(PmFinding("invalid_dependency", f"invalid dependency for {item_id}", severity="blocking", item_id=item_id))
            continue
        dependency_kind = cast(Literal["requires_code", "order_only"], kind)
        candidate = WbsDependency(assigned.get(predecessor, predecessor), dependency_kind, _none(row.get("Justification")))
        values = [value for value in values if not (value.task_id == candidate.task_id and value.kind == candidate.kind)]
        if action != "Delete":
            values.append(candidate)
    return tuple(values)


def _trace_rows(
    sheet: Any, max_text_chars: int
) -> dict[str, tuple[dict[str, Any], ...]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in _sheet_records(sheet, max_text_chars=max_text_chars):
        task_id = str(row.get("Task ID") or "").strip()
        if not task_id:
            raise PmError("unsafe_workbook", "traceability row has no Task ID")
        grouped.setdefault(task_id, []).append(row)
    return {key: tuple(value) for key, value in grouped.items()}


def _resolve_traceability(
    rows: tuple[dict[str, Any], ...], findings: list[PmFinding], item_id: str,
    baseline: WbsTask | None,
) -> dict[str, tuple[str, ...]]:
    values: dict[str, list[str]] = {
        "requirement_refs": list(baseline.requirement_refs if baseline else ()),
        "storyboard_refs": list(baseline.storyboard_refs if baseline else ()),
        "design_refs": list(baseline.design_refs if baseline else ()),
    }
    homes = {
        "requirement": "requirement_refs",
        "scenario": "requirement_refs",
        "storyboard": "storyboard_refs",
        "design": "design_refs",
        "api": "design_refs",
    }
    for row in rows:
        kind = str(row.get("Reference Type") or "").strip().lower()
        reference = str(row.get("Reference ID") or "").strip()
        field = homes.get(kind)
        action = str(row.get("Action") or "Keep").strip()
        if field is None or not reference or action not in {"Keep", "Add", "Delete"}:
            findings.append(PmFinding("invalid_traceability", f"invalid traceability for {item_id}", severity="blocking", item_id=item_id))
            continue
        values[field] = [value for value in values[field] if value != reference]
        if action != "Delete":
            values[field].append(reference)
    return {field: tuple(items) for field, items in values.items()}


def _delete_scope_changed(
    task_id: str, base: dict[str, WbsTask], current: dict[str, WbsTask]
) -> bool:
    affected = {
        item.task_id
        for item in (*base.values(), *current.values())
        if item.task_id == task_id
        or item.parent_id == task_id
        or any(dep.task_id == task_id for dep in item.dependencies)
    }
    return any(base.get(item_id) != current.get(item_id) for item_id in affected)


def _task_snapshot(task: WbsTask | None) -> dict[str, Any] | None:
    if task is None:
        return None
    return asdict(task)


def _lines(value: Any) -> tuple[str, ...]:
    return tuple(line.strip() for line in str(value or "").splitlines() if line.strip())


def _permanent_id(change_id: str, temporary: str, group: bool = False) -> str:
    digest = hashlib.sha256(f"{change_id}:{temporary}".encode()).hexdigest()[:12].upper()
    return f"{'G' if group else 'T'}-{digest}"


def _csv_safe(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIX) else value


def _none(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _integer(value: Any, default: int) -> int:
    if isinstance(value, bool):
        return default
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return default


def _hash(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()
